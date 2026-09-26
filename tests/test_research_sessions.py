import json
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError
from typer.testing import CliRunner

from valuationagent.api.main import create_app
from valuationagent.application.research import ResearchService
from valuationagent.application.research_export import build_research_export
from valuationagent.application.research_valuation import (
    ResearchValuationAssembler,
    _period,
)
from valuationagent.cli.main import app, interactive
from valuationagent.core.documents import parse_document
from valuationagent.core.tools import NoArguments, ToolSpec
from valuationagent.llm.client import LlmError, OpenAICompatibleClient
from valuationagent.schemas.agent import SearchQuery, SearchResult
from valuationagent.schemas.models import ModelConnectionInput
from valuationagent.schemas.research import (
    DocumentSummary,
    FactCandidate,
    ResearchChoice,
    ResearchQuestion,
    ResearchTurn,
)
from valuationagent.storage.sqlite import SQLiteRunStore


class ScriptedModel:
    def __init__(self, actions):
        self.actions = iter(actions)
        self.calls = []
        self.kwargs = []

    def chat(self, messages, **kwargs):
        self.calls.append(json.loads(json.dumps(messages)))
        self.kwargs.append(kwargs)
        name, args = next(self.actions)
        return {"tool_calls": [{"id": f"call_{len(self.calls)}", "type": "function",
            "function": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}],
            "reasoning_content": "PRIVATE_TEST_REASONING"}


@pytest.fixture
def service(tmp_path):
    return ResearchService(SQLiteRunStore(tmp_path))


def candidate(file_id, **changes):
    return {"metric": "revenue", "raw_value": "12000", "unit": "万元", "period": "2025",
        "scope": "consolidated", "role": "historical", "block_id": file_id + ":1",
        "quote": "2025 年合并报表，单位万元。营业收入：12000。", **changes}


def upload(store):
    return store.save_upload("年报摘录.txt", "historical_financials", "text/plain",
        "2025 年合并报表，单位万元。营业收入：12000。".encode())["file_id"]


def ready_online_session(service, *, llm=None):
    session = service.create(llm=llm)
    session.draft.company = "测试股份"
    session.draft.ticker = "600000.SH"
    session.draft.valuation_date = date(2026, 9, 24)
    session.draft.methods = ["dcf", "pe", "ev_ebitda"]
    session.data_source_preference = "online"
    service.store.save_research(session)
    return session


def test_explicit_valuation_request_returns_pipeline_action_without_llm(service):
    session = ready_online_session(service)
    state = service.turn(session.session_id, ResearchTurn(content="请立即开始正式估值"))

    assert state["action"] == {"type": "submit_valuation"}
    assert state["session"]["status"] == "ready_for_valuation"
    assert "确定性流水线" in state["messages"][-1]["content"]
    assert any(event["tool"] == "request_formal_valuation" for event in state["events"])


def test_model_can_request_formal_valuation_as_a_typed_terminal_action(service):
    model = ScriptedModel([("request_formal_valuation", {})])
    session = ready_online_session(service, llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="按刚才确认的方案继续下一步"))

    assert state["action"] == {"type": "submit_valuation"}
    assert "request_formal_valuation" in model.calls[0][0]["content"]
    assert any(event["tool"] == "request_formal_valuation" for event in state["events"])


def test_valuation_request_with_missing_scope_returns_gaps_instead_of_fake_command(service):
    session = service.create()
    state = service.turn(session.session_id, ResearchTurn(content="/valuation"))

    assert "action" not in state
    assert "研究公司尚未确认" in state["messages"][-1]["content"]
    assert "尚未提交计算" in state["messages"][-1]["content"]
    assert "单独" not in state["messages"][-1]["content"]


def test_incomplete_web_valuation_uses_model_instead_of_stopping_at_gap_list(service):
    # This test exercises a successful confirmation, so its source must also
    # identify the confirmed issuer rather than relying on a warning-only card.
    fid = service.store.save_upload("年报摘录.txt", "historical_financials", "text/plain",
        "测试股份600000.SH\n2025 年合并报表，单位万元。营业收入：12000。".encode())["file_id"]
    model = ScriptedModel([
        ("finish_response", {"answer": "还缺很多字段，请补齐后再开始估值。"}),
        ("propose_facts", {"candidates": [candidate(fid)]}),
        ("finish_response", {"answer": "收入已核验并暂存；当前摘录没有普通股数，不能补零。",
            "question": "股数不在当前摘录，如何补证？", "options": ["提供完整年报", "补充股本变动公告"]}),
    ])
    session = service.create(llm=model)
    session.draft = session.draft.model_copy(update={
        "company": "测试股份",
        "ticker": "600000.SH",
        "industry": "电子",
        "valuation_date": date(2026, 9, 24),
        "methods": ["dcf", "pe"],
    })
    session.data_source_preference = "web"
    service.store.save_research(session)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="开始正式估值", file_ids=[fid]),
    )

    assert state["session"]["pending_action"] == "valuation"
    assert state["session"]["question"]["kind"] == "clarification"
    assert state["session"]["facts"][0]["metric"] == "revenue"
    assert state["session"]["facts"][0]["status"] == "proposed"
    assert len(model.calls) == 3
    assert any(
        event["tool"] == "finish_response" and event["status"] == "failed"
        for event in state["events"]
    )
    registered = [tool["function"]["name"] for tool in model.kwargs[0]["tools"]]
    assert "search_sources" in registered
    assert "propose_facts" in registered
    assert "request_formal_valuation" in registered
    assert "不要以准备缺口清单结束" in model.calls[0][-1]["content"]


def test_confirming_facts_resumes_pending_valuation_research_automatically(service):
    text = (
        "测试股份600000.SH\n2024 年合并报表，单位万元。营业收入：12000。"
        "归属于母公司股东的净利润：1800。"
    )
    meta = service.store.save_upload(
        "年报摘录.txt", "historical_financials", "text/plain", text.encode()
    )
    fid = meta["file_id"]

    model = ScriptedModel([("propose_facts", {"candidates": [candidate(
        fid,
        metric="net_income_parent",
        raw_value="1800",
        quote="归属于母公司股东的净利润：1800。",
        period="2024",
    )]}), ("finish_response", {"answer": "已继续提取净利润，尚未确认新值。",
        "question": "摘录缺少现金和债务明细，如何补证？", "options": ["提供完整年报", "提供财务附注"]})])
    session = service.create(llm=model)
    service.store.save_research_blocks(session.session_id, fid, [{
        "block_id": fid + ":1",
        "text": text,
        "location": {"paragraph": 1, "source_type": "uploaded_document"},
    }])
    session.documents.append(DocumentSummary(
        file_id=fid, name="年报摘录.txt", role="historical_financials",
        block_count=1, sha256=meta["sha256"], size_bytes=meta["size_bytes"],
    ))
    session.draft = session.draft.model_copy(update={
        "company": "测试股份", "ticker": "600000.SH", "industry": "电子",
        "valuation_date": date(2026, 9, 24), "methods": ["dcf"],
    })
    session.data_source_preference = "web"
    session.pending_action = "valuation"
    first = FactCandidate(
        fact_id="fact_revenue", metric="revenue", raw_value="12000", unit="万元",
        normalized_value="120000000", period="2024", scope="consolidated",
        block_id=fid + ":1", quote="营业收入：12000。",
    )
    session.facts.append(first)
    session.question = ResearchQuestion(
        question_id="question_revenue", kind="facts", title="确认营业收入",
        options=[
            ResearchChoice(id="accept", label="确认"),
            ResearchChoice(id="reject", label="拒绝"),
        ],
        fact_ids=[first.fact_id],
    )
    session.status = "waiting_confirmation"
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(
        question_id="question_revenue", option_id="accept",
    ))

    assert len(model.calls) == 2
    assert state["session"]["pending_action"] == "valuation"
    assert state["session"]["facts"][0]["status"] == "confirmed"
    assert state["session"]["facts"][1]["metric"] == "net_income_parent"
    assert state["session"]["facts"][1]["status"] == "proposed"
    assert state["session"]["question"]["kind"] == "clarification"
    assert "本批实际确认 1 个字段" in state["messages"][-1]["content"]
    assert '"pending_action":"valuation"' in model.calls[0][0]["content"]


def test_final_fact_confirmation_submits_pending_valuation_without_second_start(service):
    session = service.create()
    session.draft = session.draft.model_copy(update={
        "company": "测试股份", "ticker": "600000.SH", "industry": "电子",
        "valuation_date": date(2026, 9, 24), "methods": ["dcf"],
    })
    session.data_source_preference = "web"
    session.pending_action = "valuation"
    values = {
        "revenue": ("1000000000", "元"),
        "ebit_margin": ("0.10", "ratio"),
        "tax_rate": ("0.15", "ratio"),
        "depreciation_amortization": ("50000000", "元"),
        "capital_expenditure": ("80000000", "元"),
        "change_operating_nwc": ("20000000", "元"),
        "cash_and_non_operating_assets": ("300000000", "元"),
        "interest_bearing_debt": ("200000000", "元"),
        "common_shares": ("100000000", "股"),
        "net_income_parent": ("90000000", "元"),
        "ebitda": ("150000000", "元"),
    }
    for year in range(2021, 2025):
        source_text = f"测试股份600000.SH {year}年合并报表，单位：元。\n" + "\n".join(
            f"{metric}（{unit}） {value}" for metric, (value, unit) in values.items()
        )
        meta = service.store.save_upload(f"synthetic-{year}.txt", "historical_financials", "text/plain", source_text.encode())
        blocks, _ = parse_document(service.store.get_file(meta["file_id"]))
        service.store.save_research_blocks(session.session_id, meta["file_id"], blocks)
        session.documents.append(DocumentSummary(file_id=meta["file_id"], name=f"synthetic-{year}.txt",
            role="historical_financials", block_count=len(blocks), sha256=meta["sha256"]))
        for metric, (value, unit) in values.items():
            session.facts.append(FactCandidate(
                fact_id=f"fact_{metric}_{year}",
                metric=metric,
                raw_value=value,
                unit=unit,
                normalized_value=value,
                period=str(year),
                scope="consolidated",
                block_id=blocks[0]["block_id"],
                quote=f"{metric}（{unit}） {value}",
            ))
    session.question = ResearchQuestion(
        question_id="question_snapshot", kind="facts", title="确认完整年度快照",
        options=[
            ResearchChoice(id="accept", label="确认"),
            ResearchChoice(id="reject", label="拒绝"),
        ],
        fact_ids=[fact.fact_id for fact in session.facts],
    )
    session.status = "waiting_confirmation"
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(
        question_id="question_snapshot", option_id="accept",
    ))

    assert state["action"] == {"type": "submit_valuation"}
    assert state["session"]["status"] == "ready_for_valuation"
    assert all(fact["status"] == "confirmed" for fact in state["session"]["facts"])
    assert "正在交给确定性流水线" in state["messages"][-1]["content"]


def test_preparation_prioritizes_latest_partial_year_instead_of_dumping_every_year(service):
    session = service.create()
    session.draft = session.draft.model_copy(update={
        "company": "比亚迪", "ticker": "002594.SZ", "industry": "汽车",
        "valuation_date": date(2025, 3, 31), "methods": ["dcf", "pe"],
    })
    session.data_source_preference = "web"
    for year in (2022, 2023, 2024):
        for metric, value in (("revenue", "1000000000"), ("net_income_parent", "90000000")):
            session.facts.append(FactCandidate(
                fact_id=f"fact_{metric}_{year}", metric=metric, raw_value=value,
                unit="元", normalized_value=value, period=str(year), scope="consolidated",
                block_id="message:test", quote=f"{year} {metric} {value}", status="confirmed",
            ))
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(content="/prepare"))
    answer = state["messages"][-1]["content"]

    assert "优先补齐最近年度 2024-12-31" in answer
    assert "另有 2 个不完整比较期" in answer
    assert "2022-12-31:" not in answer


def _confirmed_statement_fact(metric, value, *, period="2024", unit="元", suffix=""):
    return FactCandidate(
        fact_id=f"fact_{metric}_{period}_{suffix or 'base'}",
        metric=metric,
        raw_value=str(value),
        unit=unit,
        normalized_value=str(value),
        period=period,
        scope="consolidated",
        block_id=f"message:{metric}_{period}_{suffix or 'base'}",
        quote=f"{metric} {value}",
        status="confirmed",
    )


def test_financial_period_parser_distinguishes_annual_ranges_and_interims():
    assert _period("2024年1-12月") == date(2024, 12, 31)
    assert _period("2024-01-01 至 2024-12-31") == date(2024, 12, 31)
    assert _period("2024年度") == date(2024, 12, 31)
    assert _period("2024年第三季度") is None
    assert _period("2024H1") is None


def test_structured_assembler_deterministically_derives_annual_report_metrics(service):
    session = service.create()
    raw_values = {
        "revenue": "1000",
        "profit_before_tax": "100",
        "income_tax_expense": "15",
        "interest_expense": "10",
        "depreciation_fixed_assets": "20",
        "amortization_intangible_assets": "5",
        "amortization_long_term_deferred_expenses": "2",
        # Synthetic source explicitly discloses zero ROU amortization; a
        # missing value must not be interpreted as zero for this lease balance.
        "depreciation_right_of_use": "0",
        "cash_paid_for_ppe_intangibles": "40",
        "inventory_decrease": "-2",
        "operating_receivables_decrease": "-3",
        "operating_payables_increase": "4",
        "cash_and_non_operating_assets": "200",
        "short_term_borrowings": "50",
        "current_portion_non_current_liabilities": "10",
        "long_term_borrowings": "60",
        "bonds_payable": "20",
        "lease_liabilities": "5",
        "common_shares": "100",
        "net_income_parent": "80",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit="股" if metric == "common_shares" else "元",
        )
        for metric, value in raw_values.items()
    ]

    snapshot = ResearchValuationAssembler()._structured_financials(session)[0]

    assert snapshot.ebit_margin == Decimal("0.11")
    assert snapshot.tax_rate == Decimal("0.15")
    assert snapshot.depreciation_amortization == Decimal(27)
    assert snapshot.capital_expenditure == Decimal(40)
    assert snapshot.change_operating_nwc == Decimal(1)
    assert snapshot.interest_bearing_debt == Decimal(145)
    assert snapshot.ebitda == Decimal(137)
    assert snapshot.calculation_methods["tax_rate"] == (
        "income_tax_expense / profit_before_tax"
    )
    assert snapshot.calculation_methods["change_operating_nwc"].startswith("-(")
    assert snapshot.statement_items["profit_before_tax"] == Decimal(100)
    assert {
        evidence.evidence_id for evidence in snapshot.evidence["ebitda"]
    } == {
        "fact_profit_before_tax_2024_base",
        "fact_interest_expense_2024_base",
        "fact_depreciation_fixed_assets_2024_base",
        "fact_amortization_intangible_assets_2024_base",
        "fact_amortization_long_term_deferred_expenses_2024_base",
        "fact_depreciation_right_of_use_2024_base",
    }
    assert "确定性推导" in snapshot.source_label


def test_structured_assembler_blocks_conflicting_confirmed_aliases(service):
    session = service.create()
    session.facts = [
        _confirmed_statement_fact("revenue", "1000", suffix="a"),
        _confirmed_statement_fact("营业收入", "2000", suffix="b"),
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "同期间同口径冲突" in error
    assert "未静默覆盖" in error
    assert "1000 与 2000" in error


def test_structured_assembler_does_not_clamp_anomalous_effective_tax_rate(service):
    session = service.create()
    required_direct = {
        "revenue": "1000",
        "ebit_margin": "0.1",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
        "profit_before_tax": "100",
        "income_tax_expense": "80",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=("ratio" if metric == "ebit_margin" else "股" if metric == "common_shares" else "元"),
        )
        for metric, value in required_direct.items()
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "所得税率" in error
    assert "0.6" not in error


def test_structured_assembler_records_passed_financial_reconciliations(service):
    session = service.create()
    values = {
        "revenue": "1000",
        "ebit": "100",
        "ebit_margin": "0.1",
        "profit_before_tax": "100",
        "income_tax_expense": "15",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in values.items()
    ]

    snapshot = ResearchValuationAssembler()._structured_financials(session)[0]

    assert snapshot.calculation_methods["reconciliation.ebit_margin"].startswith("passed:")
    assert snapshot.calculation_methods["reconciliation.tax_rate"].startswith("passed:")
    assert snapshot.calculation_methods["reconciliation.ebitda"].startswith("passed:")


def test_structured_assembler_blocks_financial_reconciliation_conflicts(service):
    session = service.create()
    session.facts = [
        _confirmed_statement_fact("profit_before_tax", "100"),
        _confirmed_statement_fact("income_tax_expense", "30"),
        _confirmed_statement_fact("tax_rate", "0.15", unit="ratio"),
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "财务勾稽冲突" in error
    assert "所得税率直接值为 0.15" in error
    assert "复算为 0.3" in error


def test_one_complete_year_is_blocked_before_formal_model_handoff(service):
    session = service.create()
    session.draft = session.draft.model_copy(update={
        "company": "测试汽车",
        "industry": "汽车制造",
        "valuation_date": date(2025, 4, 30),
        "methods": ["dcf"],
    })
    session.data_source_preference = "web"
    values = {
        "revenue": "1000",
        "ebit_margin": "0.1",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in values.items()
    ]

    assembler = ResearchValuationAssembler()
    error = assembler.structured_readiness_error(session)

    assert "至少需要4个连续年度完整快照" in error
    assert "当前只有 1 个完整年度（2024）" in error
    assert "2021、2022、2023" in error
    with pytest.raises(ValueError, match="至少需要4个连续年度"):
        assembler.build(session)


def test_newer_partial_year_is_not_hidden_by_an_older_complete_snapshot(service):
    session = service.create()
    complete_values = {
        "revenue": "900",
        "ebit_margin": "0.1",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            period="2023",
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in complete_values.items()
    ]
    session.facts.append(
        _confirmed_statement_fact("revenue", "1000", period="2024")
    )

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "优先补齐最近年度 2024-12-31" in error
    assert "归母净利润" in error


def test_initial_ticker_request_becomes_confirmable_scope_before_model_research(service):
    model = ScriptedModel([("finish_response", {"answer": "不应在确认范围前调用"})])
    session = service.create(llm=model)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="我想研究 600276 恒瑞医药的历史财务数据并估值"),
    )

    proposed = state["session"]["question"]["proposed_draft"]
    assert state["session"]["question"]["kind"] == "task"
    assert proposed["ticker"] == "600276"
    assert proposed["company"] == "恒瑞医药"
    assert proposed["valuation_date"]
    assert proposed["methods"] == ["dcf", "pe", "ev_ebitda"]
    assert state["session"]["draft"]["ticker"] == ""
    assert model.calls == []


def test_initial_ticker_request_preserves_explicit_date_and_methods(service):
    model = ScriptedModel([("finish_response", {"answer": "不应在确认范围前调用"})])
    session = service.create(llm=model)

    proposed_state = service.turn(
        session.session_id,
        ResearchTurn(
            content="研究 600276 恒瑞医药，估值日：2025-12-31，方法：DCF、PE"
        ),
    )
    proposed = proposed_state["session"]["question"]["proposed_draft"]
    assert proposed["valuation_date"] == "2025-12-31"
    assert proposed["methods"] == ["dcf", "pe"]

    accepted = service.turn(
        session.session_id,
        ResearchTurn(
            question_id=proposed_state["session"]["question"]["question_id"],
            option_id="accept",
        ),
    )
    assert accepted["session"]["draft"]["valuation_date"] == "2025-12-31"
    assert accepted["session"]["draft"]["methods"] == ["dcf", "pe"]
    assert model.calls == []


def test_search_snippet_financial_fact_requires_repair_not_blanket_confirmation(service):
    text = "2023 年合并报表，单位亿元。营业收入：12000。"
    file_id = "web_search_lead"
    model = ScriptedModel([("propose_facts", {"candidates": [candidate(
        file_id,
        unit="亿元",
        period="2023",
        quote=text,
    )]})] * 3)
    session = service.create(llm=model)
    service.store.save_research_blocks(session.session_id, file_id, [{
        "block_id": file_id + ":1",
        "text": text,
        "location": {
            "source_type": "web_search",
            "url": "https://static.cninfo.com.cn/example.pdf",
            "provider": "test",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=file_id,
        name="搜索摘要",
        role="evidence",
        block_count=1,
        warnings=["联网搜索摘要；形成关键事实前应打开原始URL核对全文。"],
    ))
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(content="提取这个搜索结果"))

    assert "搜索摘要" in state["session"]["facts"][0]["warnings"][0]
    assert len(model.calls) == 3
    assert len(state["session"]["facts"]) == 1
    assert state["session"]["last_issue"]["code"] == "EVIDENCE_REPAIR_EXHAUSTED"
    assert [option["id"] for option in state["session"]["question"]["options"]] == [
        "retry", "revise", "defer",
    ]


def test_explicit_small_field_list_is_batched_into_one_confirmation(service):
    text = "2024 年合并报表，单位万元。营业收入：12000。归母净利润：1800。"
    meta = service.store.save_upload(
        "关键财务.txt", "historical_financials", "text/plain", text.encode()
    )
    file_id = meta["file_id"]
    revenue = candidate(
        file_id,
        period="2024",
        quote="营业收入：12000。",
    )
    profit = candidate(
        file_id,
        metric="net_income_parent",
        raw_value="1800",
        period="2024",
        quote="归母净利润：1800。",
    )
    model = ScriptedModel([
        ("propose_facts", {"candidates": [profit]}),
        ("propose_facts", {"candidates": [revenue, profit]}),
    ])
    session = service.create(llm=model)
    service.store.save_research_blocks(session.session_id, file_id, [{
        "block_id": file_id + ":1",
        "text": text,
        "location": {"paragraph": 1, "source_type": "uploaded_document"},
    }])
    session.documents.append(DocumentSummary(
        file_id=file_id,
        name="关键财务.txt",
        role="historical_financials",
        block_count=1,
        sha256=meta["sha256"],
        size_bytes=meta["size_bytes"],
    ))
    service.store.save_research(session)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="请提取营业收入和归母净利润，合并给我确认。"),
    )

    assert len(model.calls) == 2
    assert {fact["metric"] for fact in state["session"]["facts"]} == {
        "revenue", "net_income_parent",
    }
    assert state["session"]["question"]["kind"] == "facts"
    assert len(state["session"]["question"]["fact_ids"]) == 2
    fact_tool_events = [
        event for event in state["events"]
        if event["tool"] == "propose_facts"
    ]
    assert sum(event["status"] == "completed" for event in fact_tool_events) == 2
    assert not any(event["status"] == "failed" for event in fact_tool_events)


def test_repeated_partial_fact_batch_preserves_valid_candidate_without_recovery(service):
    text = "2024 年合并报表，单位万元。归母净利润：1800。"
    meta = service.store.save_upload(
        "关键财务.txt", "historical_financials", "text/plain", text.encode()
    )
    file_id = meta["file_id"]
    profit = candidate(
        file_id,
        metric="net_income_parent",
        raw_value="1800",
        period="2024",
        quote="归母净利润：1800。",
    )
    model = ScriptedModel([
        ("propose_facts", {"candidates": [profit]}),
        ("propose_facts", {"candidates": [profit]}),
        ("propose_facts", {"candidates": [profit]}),
    ])
    session = service.create(llm=model)
    service.store.save_research_blocks(session.session_id, file_id, [{
        "block_id": file_id + ":1",
        "text": text,
        "location": {"paragraph": 1, "source_type": "uploaded_document"},
    }])
    session.documents.append(DocumentSummary(
        file_id=file_id,
        name="关键财务.txt",
        role="historical_financials",
        block_count=1,
        sha256=meta["sha256"],
        size_bytes=meta["size_bytes"],
    ))
    service.store.save_research(session)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="请提取营业收入和归母净利润，合并给我确认。"),
    )

    assert len(model.calls) == 3
    assert state["session"]["question"]["kind"] == "facts"
    assert [fact["metric"] for fact in state["session"]["facts"]] == [
        "net_income_parent"
    ]
    assert any("仍未形成可靠候选：营业收入" in gap for gap in state["session"]["gaps"])
    assert state["session"]["last_issue"] is None


def test_official_search_pdf_can_be_downloaded_parsed_and_traced(tmp_path):
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "2023 consolidated revenue 12000 CNY")
    pdf.save()

    def handler(request):
        assert request.url.host == "static.cninfo.com.cn"
        return httpx.Response(
            200,
            headers={"content-type": "application/pdf"},
            content=output.getvalue(),
        )

    service = ResearchService(
        SQLiteRunStore(tmp_path / "remote-pdf"),
        remote_transport=httpx.MockTransport(handler),
    )
    session = service.create()
    lead_id = "web_cninfo_report"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "annual report",
        "location": {
            "source_type": "web_search",
            "url": "https://static.cninfo.com.cn/finalpage/report.pdf",
            "provider": "mock-search",
            "search_query": "annual report",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Annual report lead",
        role="evidence",
        block_count=1,
        warnings=["联网搜索摘要；形成关键事实前应打开原始URL核对全文。"],
    ))

    result = service._fetch_search_source(session, lead_id)
    blocks = service.store.research_blocks(session.session_id, result["file_id"])

    assert result["status"] == "fetched"
    assert "revenue 12000" in blocks[0]["text"]
    assert blocks[0]["location"]["source_type"] == "remote_document"
    assert blocks[0]["location"]["parent_search_file_id"] == lead_id
    assert blocks[0]["location"]["source_url"].startswith("https://static.cninfo.com.cn/")


def test_public_search_html_can_be_downloaded_read_and_traced(tmp_path):
    page = b"""<!doctype html><html><head><style>hidden</style></head><body>
    <h1>Industry policy</h1><p>Effective from 2025-01-01.</p>
    <table><tr><th>Measure</th><th>Value</th></tr><tr><td>Subsidy</td><td>10%</td></tr></table>
    <script>stealSecrets()</script></body></html>"""

    def handler(request):
        assert request.url == httpx.URL("https://example.com/policy")
        return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, content=page)

    service = ResearchService(
        SQLiteRunStore(tmp_path / "remote-html"),
        remote_transport=httpx.MockTransport(handler),
    )
    session = service.create()
    lead_id = "web_public_policy"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "policy lead",
        "location": {
            "source_type": "web_search",
            "url": "https://example.com/policy",
            "provider": "tavily",
            "search_query": "industry policy",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Policy lead",
        role="evidence",
        block_count=1,
    ))

    result = service._fetch_search_source(session, lead_id)
    blocks = service.store.research_blocks(session.session_id, result["file_id"])
    extracted = "\n".join(block["text"] for block in blocks)

    assert result["status"] == "fetched"
    assert "Industry policy" in extracted
    assert "Subsidy" in extracted and "10%" in extracted
    assert "stealSecrets" not in extracted and "hidden" not in extracted
    assert blocks[0]["location"]["source_type"] == "remote_web_document"
    assert blocks[0]["location"]["parent_search_file_id"] == lead_id
    assert any("公开网页原文" in warning for warning in result["warnings"])


def test_public_search_fetch_rejects_private_network_targets(tmp_path):
    service = ResearchService(
        SQLiteRunStore(tmp_path / "blocked-web"),
        remote_transport=httpx.MockTransport(lambda request: httpx.Response(200)),
    )
    session = service.create()
    lead_id = "web_private_target"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "unsafe lead",
        "location": {
            "source_type": "web_search",
            "url": "https://127.0.0.1/internal",
            "provider": "tavily",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Unsafe lead",
        role="evidence",
        block_count=1,
    ))

    with pytest.raises(ValueError, match="私有、回环或保留地址"):
        service._fetch_search_source(session, lead_id)


def test_tushare_opt_out_stops_submission_and_requires_a_new_source_choice(service):
    model = ScriptedModel([("request_formal_valuation", {})])
    session = ready_online_session(service, llm=model)

    state = service.turn(session.session_id, ResearchTurn(content="不使用 Tushare"))

    assert "action" not in state
    assert model.calls == []
    assert state["session"]["data_source_preference"] == ""
    assert state["session"]["question"]["kind"] == "data_source"
    assert [option["id"] for option in state["session"]["question"]["options"]] == [
        "web", "upload", "online",
    ]
    assert "已停止自动提交" in state["messages"][-1]["content"]

    question = state["session"]["question"]
    state = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="upload",
    ))
    assert state["session"]["data_source_preference"] == "upload"
    assert "不会用联网结果替代" in state["messages"][-1]["content"]


def test_model_can_propose_a_data_source_change_without_submitting(service):
    model = ScriptedModel([("propose_data_source_change", {
        "source": "upload", "reason": "用户希望改用自己准备的审计报表",
    })])
    session = ready_online_session(service, llm=model)

    state = service.turn(session.session_id, ResearchTurn(content="改用我自己的资料"))

    assert "action" not in state
    assert state["session"]["question"]["kind"] == "data_source"
    assert any(event["tool"] == "propose_data_source_change" for event in state["events"])


def test_tushare_opt_out_can_continue_with_web_research(service):
    session = ready_online_session(service)
    state = service.turn(session.session_id, ResearchTurn(content="不使用 Tushare"))
    question = state["session"]["question"]

    state = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="web",
    ))

    assert state["session"]["data_source_preference"] == "web"
    assert "公开资料自动检索" in state["messages"][-1]["content"]
    assert "关键数据不可得" in state["messages"][-1]["content"]

    state = service.turn(session.session_id, ResearchTurn(content="/prepare"))
    assert "已选择联网检索" in state["messages"][-1]["content"]
    assert "尚无可用于估值的已确认财务字段" in state["messages"][-1]["content"]


def test_confirmed_web_source_starts_search_without_reopening_source_question(service):
    model = ScriptedModel([("search_sources", {
        "query": "600276 恒瑞医药 公司业务简介",
        "reason": "检索公司业务背景",
        "purpose": "company_profile",
    })])
    session = ready_online_session(service, llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="不使用 Tushare"))
    question = state["session"]["question"]

    state = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="web",
    ))

    assert state["session"]["data_source_preference"] == "web"
    assert state["session"]["question"]["kind"] == "search_unavailable"
    names = [tool["function"]["name"] for tool in model.kwargs[0]["tools"]]
    assert "search_sources" in names
    assert "propose_facts" not in names
    assert "propose_data_source_change" not in names
    assert "数据来源 web 已经由用户确认" in model.calls[0][-1]["content"]


def test_web_financial_search_uses_official_catalogue_without_tavily_key(tmp_path):
    class OfficialCatalogue:
        def search_annual_reports(self, ticker, years, *, cutoff, company_name):
            assert ticker == "600000.SH"
            assert years == [2022, 2023, 2024]
            return SearchResult(
                query=SearchQuery(
                    query="600000 2022 2023 2024 年度报告",
                    ticker=ticker,
                    purpose="financials",
                ),
                provider="cninfo-announcements",
                provider_version="test",
                status="completed",
                hits=[{
                    "source_id": "cninfo_report2024",
                    "title": "测试股份2024年年度报告",
                    "url": "https://static.cninfo.com.cn/finalpage/2025-03-31/report.PDF",
                    "domain": "static.cninfo.com.cn",
                    "snippet": "巨潮资讯官方公告目录；证券代码 600000；报告年度 2024。",
                    "published_at": "2025-03-31",
                    "relevance": 1,
                }],
            )

    model = ScriptedModel([
        ("search_sources", {
            "query": "600000 2024/2023/2022 年度报告",
            "reason": "补齐历史年报",
            "purpose": "financials",
        }),
        ("finish_response", {
            "answer": "已从巨潮资讯官方公告目录定位年报，下一步下载原文。",
        }),
    ])
    service = ResearchService(
        SQLiteRunStore(tmp_path / "official-catalogue"),
        official_search_provider=OfficialCatalogue(),
    )
    session = ready_online_session(service, llm=model)
    session.data_source_preference = "web"
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(content="继续查找这三年的官方年报"))

    assert state["session"]["question"] is None
    assert any(doc["file_id"] == "web_cninfo_report2024" for doc in state["session"]["documents"])
    assert "官方公告目录" in state["messages"][-1]["content"]
    assert len(model.calls) == 2


def test_candidate_input_tolerates_common_scope_labels_and_adjacent_columns(service):
    text = "合并报表，单位元。\n项目 2025年 2024年\n资本开支 1,286,898,447.55 1,550,236,879.62"
    fid = service.store.save_upload(
        "现金流.txt", "historical_financials", "text/plain", text.encode()
    )["file_id"]
    model = ScriptedModel([("propose_facts", {"candidates": [{
        "metric": "capital_expenditure",
        "raw_value": "1,286,898,447.55",
        "unit": "人民币元",
        "period": "2025年度",
        "scope": "合并（上市公司）",
        "block_id": fid + ":1",
        "quote": text,
    }]})])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取资本开支", file_ids=[fid]))
    assert state["session"]["facts"][0]["scope"] == "consolidated"
    assert state["session"]["facts"][0]["unit"] == "元"
    assert state["session"]["question"]["kind"] == "facts"


def test_failed_search_stops_before_a_second_llm_call(service):
    class FailedSearch:
        provider_id = "failed-test"
        version = "1"

        def search(self, query):
            return SearchResult(
                query=query,
                provider=self.provider_id,
                status="failed",
                error_code="SEARCH_AUTH_FAILED",
                error_message="Tavily API Key 无效或无权访问，请重新配置 Tavily Key。",
            )

    model = ScriptedModel([("search_sources", {
        "query": "测试公司 年报", "reason": "缺少年报",
    })])
    session = service.create(llm=model)
    session.data_source_preference = "online"
    service.store.save_research(session)
    service.attach_search(session.session_id, FailedSearch())
    state = service.turn(session.session_id, ResearchTurn(content="联网找年报"))
    assert len(model.calls) == 1
    assert state["session"]["question"]["kind"] == "search_failed"
    assert "没有把失败结果交给模型继续猜测" in state["messages"][-1]["content"]


def test_ambiguous_multi_value_candidate_is_isolated_without_losing_valid_fact(service):
    text = "2025 年合并报表，单位元。营业收入 12000。资本开支 100 90"
    fid = service.store.save_upload(
        "年报.txt", "historical_financials", "text/plain", text.encode()
    )["file_id"]
    model = ScriptedModel([("propose_facts", {"candidates": [
        {
            "metric": "revenue", "raw_value": "12000", "unit": "元",
            "period": "2025", "scope": "consolidated", "block_id": fid + ":1",
            "quote": "2025 年合并报表，单位元。营业收入 12000。",
        },
        {
            "metric": "capital_expenditure", "raw_value": "100 90", "unit": "元",
            "period": "unknown", "scope": "unknown", "block_id": fid + ":1",
            "quote": "资本开支 100 90",
        },
    ]})])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取", file_ids=[fid]))
    assert [fact["metric"] for fact in state["session"]["facts"]] == ["revenue"]
    assert any("包含多个数字" in gap for gap in state["session"]["gaps"])
    assert "隔离" in state["messages"][-1]["content"]


def test_online_ticker_preparation_excludes_unconfirmed_search_candidates(service):
    session = service.create()
    session.data_source_preference = "online"
    session.draft = session.draft.model_copy(update={
        "company": "测试公司", "ticker": "600519.SH",
        "valuation_date": date(2026, 9, 23), "methods": ["dcf"],
    })
    session.facts.append(FactCandidate(
        fact_id="fact_pending", metric="营业收入", raw_value="12000", unit="万元",
        normalized_value="120000000", period="unknown", scope="unknown",
        block_id="message:note", quote="营业收入 12000 万元",
        warnings=["期间待确认"],
    ))
    session.gaps = ["搜索摘要年份仍需核对"]
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(content="/prepare"))
    assert state["session"]["status"] == "ready_for_valuation"
    assert state["session"]["facts"][0]["status"] == "proposed"
    assert "不会进入正式计算" in state["messages"][-1]["content"]
    request = service.valuation_assembler.build(service.store.get_research(session.session_id))
    assert request.data_source == "ticker"
    assert request.financials is None


def test_research_starts_incomplete_and_requires_explicit_selection(service):
    session = service.create()
    result = service.turn(session.session_id, ResearchTurn(content="/company 测试公司"))
    question = result["session"]["question"]
    assert result["session"]["draft"]["company"] == ""
    result = service.turn(session.session_id, ResearchTurn(content="先不要确认，我想再看看"))
    assert result["session"]["draft"]["company"] == ""
    result = service.turn(session.session_id, ResearchTurn(question_id=question["question_id"], option_id="accept"))
    assert result["session"]["draft"]["company"] == "测试公司"
    assert result["session"]["question"]["kind"] == "data_source"
    assert [option["id"] for option in result["session"]["question"]["options"]] == [
        "online", "web", "upload",
    ]
    assert service.store.list_runs() == []
    before = len(result["messages"])
    with pytest.raises(ValueError, match="失效"):
        service.turn(session.session_id, ResearchTurn(question_id=question["question_id"], option_id="accept"))
    assert len(service.store.list_messages(session.session_id)) == before


def test_text_corrections_cannot_accept_old_values(service):
    session = service.create()
    state = service.turn(session.session_id, ResearchTurn(content="/company 旧名称"))
    with pytest.raises(ValueError, match="不能同时确认"):
        service.turn(session.session_id, ResearchTurn(question_id=state["session"]["question"]["question_id"], option_id="accept", content="其实应改为新名称"))
    assert service.store.get_research(session.session_id).draft.company == ""


def test_free_text_supersedes_the_exact_question_without_accepting_it(service):
    session = service.create()
    first = service.turn(session.session_id, ResearchTurn(content="/company 旧名称"))
    old_question = first["session"]["question"]
    state = service.turn(session.session_id, ResearchTurn(
        content="/company 新名称", question_id=old_question["question_id"]
    ))
    assert state["session"]["draft"]["company"] == ""
    assert state["session"]["question"]["question_id"] != old_question["question_id"]
    assert state["session"]["question"]["proposed_draft"]["company"] == "新名称"
    with pytest.raises(ValueError, match="失效"):
        service.turn(session.session_id, ResearchTurn(
            question_id=old_question["question_id"], option_id="accept"
        ))


def test_file_llm_tool_evidence_and_confirmation_survive_restart(service):
    fid = upload(service.store)
    model = ScriptedModel([("read_document", {"file_id": fid}), ("propose_facts", {"candidates": [candidate(fid)]})])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取营业收入", file_ids=[fid]))
    fact = state["session"]["facts"][0]
    assert fact["status"] == "proposed" and fact["normalized_value"] == "120000000"
    assert state["session"]["documents"][0]["block_count"] == 1
    assert state["session"]["data_source_preference"] == "upload"
    assert len(state["session"]["documents"][0]["sha256"]) == 64
    assert "PRIVATE_TEST_REASONING" not in json.dumps(state)
    # Reasoning survives only within the ephemeral protocol loop.
    assert any(m.get("reasoning_content") == "PRIVATE_TEST_REASONING" for m in model.calls[-1])
    qid = state["session"]["question"]["question_id"]
    fresh = ResearchService(SQLiteRunStore(service.store.data_dir))
    state = fresh.turn(session.session_id, ResearchTurn(question_id=qid, option_id="accept"))
    assert state["session"]["facts"][0]["status"] == "confirmed"
    assert fresh.store.research_blocks(session.session_id, fid)[0]["text"] == fact["quote"]
    assert "No formal valuation submitted" in build_research_export(fresh, session.session_id, "html")[0]


def test_long_conversation_keeps_explicit_durable_memory_after_restart(service):
    first = ScriptedModel([
        ("update_memory", {
            "updates": [{
                "key": "scope.statement_basis",
                "kind": "constraint",
                "content": "本任务后续分析统一采用合并报表口径；发现母公司口径时先询问用户。",
            }],
        }),
        ("finish_response", {
            "answer": "已记住本任务统一采用合并口径。",
        }),
    ])
    session = service.create(llm=first)
    service.turn(session.session_id, ResearchTurn(content="后续统一采用合并口径，遇到母公司口径先问我。"))
    for index in range(30):
        service.store.add_message(session.session_id, "user", f"临时讨论 {index}", "research")
        service.store.add_message(session.session_id, "assistant", f"临时回复 {index}", "research")

    fresh = ResearchService(SQLiteRunStore(service.store.data_dir))
    second = ScriptedModel([("finish_response", {"answer": "我会继续遵循已保存口径。"})])
    fresh.attach(session.session_id, second)
    state = fresh.turn(session.session_id, ResearchTurn(content="继续之前的任务"))

    system_prompt = second.calls[0][0]["content"]
    assert "scope.statement_basis" in system_prompt
    assert "统一采用合并报表口径" in system_prompt
    assert state["session"]["memory"][0]["source_message_id"].startswith("msg_")


def test_credentials_are_redacted_from_messages_memory_and_tool_events(service):
    secret = "sk-testsecret1234567890"
    model = ScriptedModel([("finish_response", {
        "answer": f"不会保存 {secret}",
        "memory_updates": [{
            "key": "preference.secret",
            "kind": "preference",
            "content": secret,
        }],
    })])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content=f"请记住 {secret}"))
    serialized = json.dumps(state, ensure_ascii=False)
    assert secret not in serialized
    assert "REDACTED_CREDENTIAL" in serialized
    assert state["session"]["memory"] == []
    assert any(event["type"] == "security.credential_redacted" for event in state["events"])


def test_model_failure_becomes_recoverable_user_choice(service):
    class FailingModel:
        def chat(self, messages, **kwargs):
            raise LlmError("LLM_HTTP_503: 供应商暂时不可用，请稍后恢复。")

    session = service.create(llm=FailingModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续整理资料"))
    question = state["session"]["question"]
    assert question["kind"] == "recovery"
    assert {item["id"] for item in question["options"]} >= {"retry", "revise", "defer"}
    assert state["session"]["last_issue"]["code"] == "LLM_HTTP_503"
    assert "未编造估值数字" in state["messages"][-1]["content"]
    assert any(event["type"] == "agent.recovery_required" for event in state["events"])

    service.attach(session.session_id, ScriptedModel([("finish_response", {"answer": "重试成功，已从原上下文继续。"})]))
    state = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="retry"
    ))
    assert state["session"]["question"] is None
    assert state["session"]["last_issue"]["status"] == "resolved"
    assert "重试成功" in state["messages"][-1]["content"]


def test_registered_tool_provider_joins_same_audit_loop(tmp_path):
    class FinanceProbeProvider:
        provider_id = "finance-probe"
        version = "test-1"

        def tool_specs(self, session):
            return [ToolSpec(
                "inspect_finance_contract",
                "读取未来金融插件的能力契约，不执行估值。",
                NoArguments,
                lambda _: {"available": False, "reason": "正式模型待接入"},
            )]

    model = ScriptedModel([
        ("inspect_finance_contract", {}),
        ("finish_response", {"answer": "已检查金融插件契约；正式模型仍待接入。"}),
    ])
    service = ResearchService(
        SQLiteRunStore(tmp_path / "provider-runtime"),
        tool_providers=[FinanceProbeProvider()],
    )
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="检查金融工具是否可用"))
    completed = [event for event in state["events"] if event["type"] == "tool.completed"]
    assert any(event["tool"] == "inspect_finance_contract" for event in completed)
    agent_start = next(event for event in state["events"] if event["type"] == "agent.started")
    assert agent_start["payload"]["tool_extensions"][0]["provider_id"] == "finance-probe"


def test_invented_number_is_rejected_and_tool_feedback_reaches_model(service):
    fid = upload(service.store)
    model = ScriptedModel([
        ("read_document", {"file_id": fid}),
        ("propose_facts", {"candidates": [candidate(fid, raw_value="999999")]}),
        ("finish_response", {"answer": "原文不支持这个数值，需要重新核对。"}),
    ])
    session = service.create(llm=model)
    result = service.turn(session.session_id, ResearchTurn(content="提取", file_ids=[fid]))
    assert result["session"]["facts"] == []
    assert any(e["type"] == "tool.failed" for e in result["events"])
    assert "原文不支持" in result["messages"][-1]["content"]


@pytest.mark.parametrize("sign", ["-", "−"])
def test_extraction_cannot_drop_a_negative_source_sign(service, sign):
    quote = f"2025 年合并报表，单位万元。净利润：{sign}12000。"
    fid = service.store.save_upload("亏损报表.txt", "historical_financials", "text/plain", quote.encode())["file_id"]
    model = ScriptedModel([
        ("propose_facts", {"candidates": [candidate(fid, metric="net_income", quote=quote)]}),
        ("propose_facts", {"candidates": [candidate(fid, metric="net_income", quote=quote, raw_value=f"{sign}12000")]}),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取净利润", file_ids=[fid]))
    assert len(state["session"]["facts"]) == 1
    assert state["session"]["facts"][0]["normalized_value"] == "-120000000"
    assert any(event["type"] == "tool.failed" for event in state["events"])


def test_extraction_understands_parenthetical_accounting_negatives(service):
    quote = "2024 年合并报表，单位万元。经营性应收项目的减少：(1,200)。"
    fid = service.store.save_upload(
        "现金流补充资料.txt",
        "historical_financials",
        "text/plain",
        quote.encode(),
    )["file_id"]
    model = ScriptedModel([("propose_facts", {"candidates": [candidate(
        fid,
        metric="operating_receivables_decrease",
        raw_value="(1,200)",
        quote=quote,
        period="2024",
    )]})])
    session = service.create(llm=model)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="提取经营性应收项目变动", file_ids=[fid]),
    )

    assert state["session"]["facts"][0]["raw_value"] == "(1,200)"
    assert state["session"]["facts"][0]["normalized_value"] == "-12000000"


def test_repeated_identical_tool_failure_stops_and_asks_user(service):
    fid = upload(service.store)
    invalid = {"candidates": [candidate(fid, raw_value="999999")]}
    model = ScriptedModel([
        ("propose_facts", invalid),
        ("propose_facts", invalid),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取", file_ids=[fid]))
    assert state["session"]["facts"] == []
    assert state["session"]["last_issue"]["code"] == "AGENT_NO_PROGRESS"
    assert state["session"]["question"]["kind"] == "recovery"
    assert [choice["id"] for choice in state["session"]["question"]["options"]] == [
        "retry", "revise", "defer",
    ]
    assert "最近失败原因" in state["messages"][-1]["content"]
    assert "候选数值未出现在引用原文" in state["messages"][-1]["content"]
    assert any(event["type"] == "tool.failed" for event in state["events"])


def test_uncertain_fields_are_not_blanket_confirmed(service):
    fid = upload(service.store)
    model = ScriptedModel([("propose_facts", {"candidates": [candidate(fid, unit="unknown")]})] * 3)
    session = service.create(llm=model)
    result = service.turn(session.session_id, ResearchTurn(content="提取", file_ids=[fid]))
    qid = result["session"]["question"]["question_id"]
    assert result["session"]["last_issue"]["code"] == "EVIDENCE_REPAIR_EXHAUSTED"
    assert [option["id"] for option in result["session"]["question"]["options"]] == [
        "retry", "revise", "defer",
    ]
    result = service.turn(session.session_id, ResearchTurn(question_id=qid, option_id="defer"))
    assert result["session"]["facts"][0]["status"] == "rejected"
    assert result["session"]["last_issue"]["status"] == "resolved"
    assert any(event["type"] == "facts.quarantined" for event in result["events"])


def test_pasted_correction_requires_confirmation_and_preserves_old_record(service):
    fid = upload(service.store)
    model = ScriptedModel([("propose_facts", {"candidates": [candidate(fid)]})])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="提取", file_ids=[fid]))
    state = service.turn(session.session_id, ResearchTurn(question_id=state["session"]["question"]["question_id"], option_id="accept"))
    old_id = state["session"]["facts"][0]["fact_id"]
    note = "我更正数据：2025 年合并报表，单位万元。营业收入：13000。"
    service.store.add_message(session.session_id, "user", note, "research")
    # The same message IDs are exposed to the model through inspect_context.
    block_id = next(b["block_id"] for b in service._blocks(session).values() if b["text"] == note)
    correction = candidate(fid, raw_value="13000", block_id=block_id, quote=note)
    service.attach(session.session_id, ScriptedModel([
        ("inspect_context", {}),
        ("propose_facts", {"candidates": [correction], "replaces": [old_id]}),
    ]))
    state = service.turn(session.session_id, ResearchTurn(content="请按上述补充更正"))
    assert [f["status"] for f in state["session"]["facts"]] == ["confirmed", "proposed"]
    assert state["session"]["facts"][1]["source_type"] == "user_note"
    state = service.turn(session.session_id, ResearchTurn(question_id=state["session"]["question"]["question_id"], option_id="accept"))
    assert [f["status"] for f in state["session"]["facts"]] == ["rejected", "confirmed"]
    assert state["session"]["facts"][0]["raw_value"] == "12000"


def test_partial_scope_update_preserves_other_fields_and_refreshes_gaps(service):
    session = service.create()
    state = service.turn(session.session_id, ResearchTurn(content="/company 测试公司"))
    service.turn(session.session_id, ResearchTurn(question_id=state["session"]["question"]["question_id"], option_id="accept"))
    service.attach(session.session_id, ScriptedModel([("propose_task", {"draft": {"valuation_date": "2026-09-14"}})]))
    state = service.turn(session.session_id, ResearchTurn(content="估值日用今天"))
    assert state["session"]["question"]["proposed_draft"]["company"] == "测试公司"
    service.turn(session.session_id, ResearchTurn(question_id=state["session"]["question"]["question_id"], option_id="accept"))
    state = service.turn(session.session_id, ResearchTurn(content="/prepare"))
    assert "估值基准日尚未确认" not in state["session"]["gaps"]
    assert "估值方法尚未确认" in state["session"]["gaps"]
    service.attach(session.session_id, ScriptedModel([
        ("update_data_gaps", {"missing": ["可比公司不足"], "reason": "重新核对当前资料"}),
        ("finish_response", {"answer": "已更新资料缺口。"}),
    ]))
    state = service.turn(session.session_id, ResearchTurn(content="重新检查缺口", language="en-US"))
    assert state["session"]["gaps"] == ["可比公司不足"]
    assert state["session"]["language"] == "en-US"


def test_search_gap_does_not_fabricate_peers(service):
    action = ("search_sources", {"query": "可比公司", "reason": "可比公司不足"})
    model = ScriptedModel([action, action])
    session = service.create(llm=model)
    result = service.turn(session.session_id, ResearchTurn(content="帮我补充同业"))
    assert result["session"]["question"]["kind"] == "data_source"
    assert result["session"]["facts"] == []
    assert "没有发出网络请求" in result["messages"][-1]["content"]

    question = result["session"]["question"]
    result = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="web"
    ))
    assert result["session"]["question"]["kind"] == "search_unavailable"
    assert result["session"]["facts"] == []
    assert "没有发出网络请求" in result["messages"][-1]["content"]
    assert "未配置 Tavily API Key" in result["messages"][-1]["content"]


def test_search_provider_can_be_attached_to_one_session_without_persisting_key(service):
    from valuationagent.search.providers import MockSearchProvider

    session = service.create()
    provider = MockSearchProvider()
    service.attach_search(session.session_id, provider)
    assert service._search_clients[session.session_id] is provider
    event = service.store.list_events(session.session_id)[-1]
    assert event.type == "search.attached"
    assert event.payload == {"provider": "mock-search", "provider_version": "0.1"}


def test_llm_receives_source_ambiguity_contract_and_free_text_guidance(service):
    model = ScriptedModel([("finish_response", {
        "answer": "表头中的 2024 与相邻页的 2023 无法可靠对应，暂未提取该列。",
        "question": "这列应按哪个报告期处理？",
        "options": ["按表头原样保留为 2024", "暂不采用这列，等待补充资料"],
    })])
    session = service.create(llm=model)
    result = service.turn(session.session_id, ResearchTurn(content="请提取这份表格"))

    system_prompt = model.calls[0][0]["content"]
    assert "不得静默修正或合理化" in system_prompt
    assert "不得把数值平移到相邻年份" in system_prompt
    assert "实际发现了什么" in system_prompt
    assert "自动提供 Chat 自由输入入口" in system_prompt
    assert result["session"]["question"]["kind"] == "clarification"
    assert [option["label"] for option in result["session"]["question"]["options"]] == [
        "按表头原样保留为 2024",
        "暂不采用这列，等待补充资料",
    ]
    assert "直接输入你的判断" not in result["messages"][-1]["content"]


def test_llm_questions_require_two_or_three_actionable_choices():
    from valuationagent.application.research import FinishResponse

    with pytest.raises(ValidationError, match="2—3"):
        FinishResponse(answer="需要确认。", question="采用哪个口径？", options=["采用历史口径"])
    assert FinishResponse(answer="当前没有需要确认的事项。").options == []


def test_selected_clarification_preserves_its_question_for_the_next_model_turn(service):
    model = ScriptedModel([
        ("finish_response", {
            "answer": "需要确认该列的报告期。",
            "question": "财务表第 C 列标为 2024，但附注称它是 2023；应采用哪个报告期？",
            "options": ["采用表头年份", "采用附注年份"],
        }),
        ("finish_response", {"answer": "将按你选择的表头年份整理候选，仍需核对确认。"}),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="这份表格的年份不一致"))
    question = state["session"]["question"]
    state = service.turn(session.session_id, ResearchTurn(
        question_id=question["question_id"], option_id="choice_0"
    ))
    prompt = model.calls[-1][0]["content"]
    assert question["title"] in prompt
    assert '"selected_option_id":"choice_0"' in prompt
    assert state["session"]["question"] is None


def test_large_fact_history_is_bounded_and_old_user_notes_are_retrievable(service):
    model = ScriptedModel([
        ("inspect_context", {"section": "facts", "query": "historical_revenue_000", "limit": 1}),
        ("inspect_context", {"section": "user_notes", "query": "早期成本口径"}),
        ("finish_response", {"answer": "已检索到早期字段和用户说明。"}),
    ])
    session = service.create(llm=model)
    session.facts = [FactCandidate(
        fact_id=f"fact_{index}", metric=f"historical_revenue_{index:03}",
        raw_value="100", unit="万元", normalized_value="1000000", period="2025",
        scope="consolidated", block_id=f"source:{index}", quote="营业收入 100 万元",
        status="confirmed",
    ) for index in range(620)]
    service.store.save_research(session)
    service.store.add_message(session.session_id, "user", "早期成本口径：统一使用主营业务成本。", "research")
    for index in range(30):
        service.store.add_message(session.session_id, "user", f"近期说明 {index}", "research")
    state = service.turn(session.session_id, ResearchTurn(content="请找回早期的字段和成本说明"))
    assert state["session"]["last_issue"] is None
    assert len(model.calls[0][0]["content"]) < 50000
    assert '"facts_omitted":' in model.calls[0][0]["content"]
    fact_output = json.loads(next(message["content"] for message in model.calls[1] if message["role"] == "tool"))
    assert fact_output["total"] == 1
    assert fact_output["items"][0]["fact_id"] == "fact_0"
    note_output = json.loads([message["content"] for message in model.calls[2] if message["role"] == "tool"][-1])
    assert note_output["items"][0]["text"] == "早期成本口径：统一使用主营业务成本。"
    assert note_output["items"][0]["block_id"].startswith("message:")


def test_context_retrieval_paginates_without_losing_history(service):
    model = ScriptedModel([
        ("inspect_context", {"section": "user_notes", "limit": 2}),
        ("inspect_context", {"section": "user_notes", "offset": 2, "limit": 2}),
        ("finish_response", {"answer": "已读取全部说明。"}),
    ])
    session = service.create(llm=model)
    for index in range(3):
        service.store.add_message(session.session_id, "user", f"历史说明 {index}", "research")
    service.turn(session.session_id, ResearchTurn(content="查看历史说明"))
    first = json.loads(next(message["content"] for message in model.calls[1] if message["role"] == "tool"))
    second = json.loads([message["content"] for message in model.calls[2] if message["role"] == "tool"][-1])
    assert first["next_offset"] == 2
    assert second["next_offset"] is None
    assert [item["text"] for item in first["items"] + second["items"]] == [
        "历史说明 0", "历史说明 1", "历史说明 2", "查看历史说明",
    ]


@pytest.mark.parametrize("response", [None, {"tool_calls": "invalid"}, {"tool_calls": [None]},
    {"tool_calls": [{"id": "call_bad", "function": None}]},
    {"tool_calls": [{"id": "call_bad", "function": {"name": [], "arguments": "{}"}}]}])
def test_malformed_model_protocol_becomes_actionable_recovery(service, response):
    class MalformedModel:
        def chat(self, messages, **kwargs):
            return response

    session = service.create(llm=MalformedModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续整理资料"))
    assert state["session"]["last_issue"]["code"] == "TOOL_RESPONSE_INVALID"
    assert state["session"]["question"]["kind"] == "recovery"
    assert not any(event["type"] == "tool.started" for event in state["events"])


def test_model_errors_do_not_persist_credentials_in_recovery_or_audit(service):
    secret = "sk-testcredential123456789"

    class FailingModel:
        def chat(self, messages, **kwargs):
            raise LlmError("LLM_HTTP_400: 网关返回错误，凭证 " + secret)

    session = service.create(llm=FailingModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续"))
    assert secret not in json.dumps(state, ensure_ascii=False)
    assert "REDACTED_CREDENTIAL" in state["session"]["last_issue"]["message"]


def test_tool_errors_and_question_labels_do_not_persist_credentials(service):
    secret = "sk-testcredential123456789"

    class FailingProvider:
        provider_id = "failing-provider"
        version = "test"

        def tool_specs(self, session):
            def fail(_):
                raise ValueError("上游拒绝凭证 " + secret)
            return [ToolSpec("failing_tool", "test", NoArguments, fail)]

    service.tool_providers = (FailingProvider(),)
    model = ScriptedModel([
        ("failing_tool", {}),
        ("finish_response", {"answer": "需要重新连接。", "question": "重新输入 " + secret,
                             "options": ["修改 " + secret, "稍后处理"]}),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="继续"))
    assert secret not in json.dumps(state, ensure_ascii=False)
    assert state["session"]["question"]["kind"] == "clarification"
    assert secret not in json.dumps(model.calls[-1], ensure_ascii=False)


def test_excel_locations_and_formulas_are_not_executed(service):
    book = Workbook()
    book.active.title = "历史财务"
    book.active.append(["营业收入", 12000, "万元"])
    book.active.append(["外部公式", '=WEBSERVICE("https://invalid.example")'])
    data = BytesIO()
    book.save(data)
    meta = service.store.save_upload("财务.xlsx", "historical_financials", None, data.getvalue())
    blocks, _ = parse_document(service.store.get_file(meta["file_id"]))
    assert blocks[0]["location"] == {"sheet": "历史财务", "row": 1}
    assert "B1: 12000" in blocks[0]["text"]
    assert "公式，未求值" in blocks[1]["text"]
    assert "https://" not in blocks[1]["text"]


def test_docx_upload_preserves_paragraph_and_table_order(service):
    from docx import Document

    document = Document()
    document.add_heading("资本开支假设", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1.0"
    document.add_paragraph("表后说明")
    data = BytesIO()
    document.save(data)

    meta = service.store.save_upload(
        "估值方案.docx", "evidence", None, data.getvalue()
    )
    blocks, warnings = parse_document(service.store.get_file(meta["file_id"]))

    assert warnings == []
    assert [block["text"] for block in blocks] == [
        "资本开支假设",
        "参数 | 数值",
        "alpha | 1.0",
        "表后说明",
    ]
    assert blocks[0]["location"]["paragraph"] == 1
    assert blocks[1]["location"] == {"table": 1, "row": 1}
    assert blocks[3]["location"]["paragraph"] == 2


def test_api_research_sources_are_scoped_and_exports_are_safe(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        first = client.post("/api/research-sessions", json={"language": "en-US"}).json()["session"]["session_id"]
        second = client.post("/api/research-sessions", json={}).json()["session"]["session_id"]
        fid = client.post("/api/files", data={"role": "evidence"}, files={"file": ("policy.txt", b'<script>alert(1)</script>', "text/plain")}).json()["file_id"]
        result = client.post(f"/api/research-sessions/{first}/messages", json={"file_ids": [fid]}).json()
        assert result["session"]["documents"][0]["name"] == "policy.txt"
        assert client.get(f"/api/research-sessions/{first}/sources/{fid}").status_code == 200
        assert client.get(f"/api/research-sessions/{second}/sources/{fid}").status_code == 404
        assert client.post(f"/api/research-sessions/{first}/messages", json={}).status_code == 422
        assert client.get(f"/api/research-sessions/{first}/export?format=json").json()["report_kind"] == "research_preparation"
        html = client.get(f"/api/research-sessions/{first}/export?format=html").text
        assert '<script>' not in html and "No formal valuation submitted" in html


def test_api_chat_submission_launches_the_deterministic_valuation_pipeline(tmp_path):
    app_instance = create_app(tmp_path)
    with TestClient(app_instance) as client:
        created = client.post("/api/research-sessions", json={}).json()
        session_id = created["session"]["session_id"]
        session = app_instance.state.store.get_research(session_id)
        session.draft.company = "测试股份"
        session.draft.ticker = "600000.SH"
        session.draft.valuation_date = date(2026, 9, 24)
        session.draft.methods = ["dcf", "pe", "ev_ebitda"]
        session.data_source_preference = "online"
        app_instance.state.store.save_research(session)

        response = client.post(
            f"/api/research-sessions/{session_id}/messages",
            json={"content": "开始正式估值"},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["action"]["type"] == "valuation_submitted"
        assert body["action"]["run_id"].startswith("run_")
        saved = app_instance.state.store.get_research(session_id)
        assert saved.valuation_run_id == body["action"]["run_id"]
        assert saved.status == "submitted"


def test_cli_opens_conversation_and_confirms_options(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("VALUATION_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("VALUATION_LLM_API_KEY", raising=False)
    monkeypatch.setattr("valuationagent.cli.main.sys.stdin.isatty", lambda: True)
    monkeypatch.setattr("valuationagent.cli.research.configure_model", lambda language: (None, ""))
    answers = iter(["/company Research Company", "/quit"])
    monkeypatch.setattr("valuationagent.cli.main.text_input", lambda *args, **kwargs: next(answers))
    def choose_for_cli(*args, **kwargs):
        values = [getattr(choice, "value", choice) for choice in args[1]]
        return "online" if "online" in values else "accept"

    monkeypatch.setattr("valuationagent.cli.main.select", choose_for_cli)
    interactive(language="en-US")
    assert "What can I do?" in capsys.readouterr().out
    assert "research" in CliRunner().invoke(app, ["--help"]).output
    assert SQLiteRunStore(tmp_path).list_research()[0].draft.company == "Research Company"


@pytest.mark.parametrize("thinking,expected", [("auto", "required"), ("enabled", "auto")])
def test_deepseek_tool_payload_and_redacted_diagnostics(thinking, expected):
    recorded = []
    secret = "TEST_KEY_NOT_REAL"

    def handler(request):
        recorded.append(json.loads(request.content))
        return httpx.Response(400, json={"error": {"message": "tool_choice rejected " + secret}})

    original = httpx.Client
    config = ModelConnectionInput(base_url="https://api.deepseek.com", model="deepseek-v4-flash", api_key=secret, thinking=thinking)
    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(handler), **kwargs
            ),
        ),
        pytest.raises(Exception) as caught,
    ):
        OpenAICompatibleClient(config).chat(
            [{"role": "user", "content": "hello"}],
            tools=[{"type": "function"}],
            tool_choice="required",
        )
    assert recorded[0]["tool_choice"] == expected
    assert "temperature" not in recorded[0]
    assert "tool_choice" in str(caught.value) and secret not in str(caught.value)


def test_llm_timeout_and_invalid_json_have_distinct_recovery_codes():
    original = httpx.Client
    config = ModelConnectionInput(
        base_url="https://example.test/v1", model="test-model", api_key="TEST_ONLY"
    )

    def timeout_handler(request):
        raise httpx.ReadTimeout("late", request=request)

    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(timeout_handler), **kwargs
            ),
        ),
        pytest.raises(LlmError, match="LLM_TIMEOUT"),
    ):
        OpenAICompatibleClient(config).chat([{"role": "user", "content": "hello"}])

    def invalid_json_handler(request):
        return httpx.Response(200, content=b"not-json")

    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(invalid_json_handler), **kwargs
            ),
        ),
        pytest.raises(LlmError, match="LLM_RESPONSE_INVALID_JSON"),
    ):
        OpenAICompatibleClient(config).chat([{"role": "user", "content": "hello"}])
