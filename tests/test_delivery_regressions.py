"""Regressions from the acceptance audit: bad evidence, recovery and scale."""
import copy
import json
import threading
from datetime import date
from decimal import Decimal as D
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from valuationagent.api.main import create_app
from valuationagent.application.research import CandidateInput, ProposeFacts, ResearchService
from valuationagent.application.research_valuation import ResearchValuationAssembler
from valuationagent.application.reproducibility import build_valuation_bundle, replay_bundle
from valuationagent.application.runner import ValuationRunner
from valuationagent.application.reporting import ValuationReportExporter
from valuationagent.core.documents import parse_document
from valuationagent.core.evidence import bind_evidence
from valuationagent.finance.team_model import FinanceTeamModel
from valuationagent.llm.client import LlmError, OpenAICompatibleClient
from valuationagent.schemas.models import FinancialSnapshot, ModelConnectionInput, PeerCompany, ValuationRequest
from valuationagent.schemas.research import DocumentSummary, FactCandidate, ResearchDraft, ResearchSession, ResearchTurn
from valuationagent.storage.sqlite import SQLiteRunStore


def source(text, **location):
    return {"block_id": "file_1:1", "file_id": "file_1", "text": text, "location": location}


def candidate(**changes):
    return CandidateInput.model_validate({"metric": "营业收入", "raw_value": "100", "unit": "万元",
        "period": "2025", "scope": "consolidated", "block_id": "file_1:1", "quote": "营业收入 | 100 | 90", **changes})


def draft():
    return ResearchDraft(company="测试股份", ticker="600123", valuation_date=date(2026, 6, 30), methods=["pe"], industry="电子")


@pytest.mark.parametrize("change,reason", [({"period": "2024"}, "年度列冲突"),
    ({"unit": "元"}, "单位冲突"), ({"scope": "parent"}, "口径冲突")])
def test_evidence_rejects_misbound_columns_units_and_scope(change, reason):
    block = source("测试股份 600123 合并利润表\n单位：万元\n项目 | 2025年 | 2024年\n营业收入 | 100 | 90")
    warnings, _ = bind_evidence(candidate(**change), block, [block], draft())
    assert any(reason in warning for warning in warnings)


def test_evidence_rejects_wrong_company_and_future_disclosure():
    block = source("其他公司 600999 合并报表\n单位：万元\n2025年\n营业收入 100", published_at="2026-07-01")
    warnings, _ = bind_evidence(candidate(), block, [block], draft())
    assert any("主体未匹配" in w for w in warnings)
    assert any("晚于估值日" in w for w in warnings)


def test_evidence_accepts_correct_table_and_preserves_source_metadata(tmp_path):
    store = SQLiteRunStore(tmp_path)
    service = ResearchService(store)
    session = service.create()
    session.draft = draft()
    text = "测试股份 600123 合并利润表\n单位：万元\n项目 | 2025年 | 2024年\n营业收入 | 100 | 90"
    meta = store.save_upload("财务.txt", "historical_financials", "text/plain", text.encode())
    block = source(text, page=24, source_url="https://example.com/report", published_at="2026-04-01")
    block.update(file_id=meta["file_id"], block_id=meta["file_id"] + ":1")
    store.save_research_blocks(session.session_id, meta["file_id"], [block])
    session.documents.append(DocumentSummary(file_id=meta["file_id"], name="财务.txt", role="historical_financials", block_count=1, sha256=meta["sha256"]))
    service._facts(session, ProposeFacts(candidates=[candidate(block_id=block["block_id"])]))
    fact = session.facts[0]
    assert not fact.warnings
    assert fact.verification["year_column"] == 2025
    ref = ResearchValuationAssembler()._evidence(session, fact)
    assert ref.page == 24 and ref.source_sha256 == meta["sha256"]
    assert ref.source_url == "https://example.com/report" and ref.published_at == date(2026, 4, 1)


def test_xlsx_formula_caches_are_streamed_not_random_access(tmp_path, monkeypatch):
    book = Workbook()
    for i in range(800):
        book.active.append([i, f"=A{i+1}*2"])
    content = BytesIO()
    book.save(content)
    store = SQLiteRunStore(tmp_path)
    meta = store.save_upload("公式.xlsx", "historical_financials", "application/octet-stream", content.getvalue())
    def forbidden(*args):
        raise AssertionError("Quadratic read-only cell access is forbidden")
    monkeypatch.setattr(ReadOnlyWorksheet, "__getitem__", forbidden)
    blocks, warnings = parse_document(store.get_file(meta["file_id"]))
    assert len(blocks) == 800
    assert any("缓存" in w for w in warnings)


def test_compact_snapshot_omits_large_payload_and_bounds_history(tmp_path):
    store = SQLiteRunStore(tmp_path)
    service = ResearchService(store)
    session = service.create()
    for i in range(250):
        store.append_event(session.session_id, type="tool.completed", stage="research", status="completed",
                           summary="done", payload={"text": "x" * 20000})
        store.add_message(session.session_id, "assistant", str(i))
    compact = service.snapshot(session.session_id, compact=True)
    assert len(compact["events"]) == 100 and len(compact["messages"]) == 60
    assert "payload" not in compact["events"][0]
    assert len(json.dumps(compact)) < 100000
    assert compact["older_messages"] and compact["messages"][-1]["content"] == "249"
    assert len(store.event_detail(session.session_id, compact["cursor"])["payload"]["text"]) == 20000
    assert not service.snapshot(session.session_id, compact=True, after=compact["cursor"])["events"]


def test_job_reservation_is_idempotent_and_conflicting_submissions_fail(tmp_path):
    service = ResearchService(SQLiteRunStore(tmp_path))
    sid = service.create().session_id
    turn = ResearchTurn(content="解释估值步骤", request_id="request_idempotent")
    assert service.reserve_turn(sid, turn) == (turn.request_id, True)
    assert service.reserve_turn(sid, turn) == (turn.request_id, False)
    with pytest.raises(ValueError, match="不同内容"):
        service.reserve_turn(sid, turn.model_copy(update={"content": "other"}))
    with pytest.raises(ValueError, match="正在执行"):
        service.reserve_turn(sid, ResearchTurn(content="another request"))


def test_cancel_queued_turn_preserves_recoverable_session(tmp_path):
    service = ResearchService(SQLiteRunStore(tmp_path))
    sid = service.create().session_id
    turn = ResearchTurn(content="解释估值步骤", request_id="request_cancelled")
    service.reserve_turn(sid, turn)
    assert service.cancel_turn(sid)["cancel_requested"]
    result = service.turn(sid, turn, reserved=True)
    assert result["execution"]["status"] == "cancelled"
    assert result["session"]["last_issue"]["code"] == "EXECUTION_CANCELLED"
    assert result["session"]["question"]["kind"] == "recovery"


def test_server_restart_marks_expired_job_interrupted(tmp_path):
    store = SQLiteRunStore(tmp_path)
    service = ResearchService(store)
    sid = service.create().session_id
    service.reserve_turn(sid, ResearchTurn(content="读取资料", request_id="request_restart"))
    with store._connect() as db:
        db.execute("UPDATE research_jobs SET updated=0 WHERE session_id=?", (sid,))
    restarted = ResearchService(store)
    assert restarted.snapshot(sid, compact=True)["execution"]["status"] == "interrupted"
    assert restarted.reserve_turn(sid, ResearchTurn(content="继续读取"))[1]


def relative_request(method="pe"):
    values = {"pe": {"net_income_parent": "200"}, "ps": {"revenue": "1000"},
              "ev_ebitda": {"ebitda": "300", "cash_and_non_operating_assets": "100", "interest_bearing_debt": "50"}}
    peers = [PeerCompany(ticker=f"P{i}", name=f"Peer {i}", **{method: str(value)}) for i, value in enumerate([10, 12, 14, 16, 18])]
    return ValuationRequest(company={"name": "测试股份", "industry": "电子"}, valuation_date="2026-06-30", methods=[method],
        financials=FinancialSnapshot(period_end="2025-12-31", common_shares="100", **values[method]), peers=peers)


@pytest.mark.parametrize("method,base,low,high", [("pe", "28", "24", "32"), ("ps", "140", "120", "160"), ("ev_ebitda", "42.5", "36.5", "48.5")])
def test_relative_only_accepts_minimal_inputs_and_matches_independent_arithmetic(tmp_path, method, base, low, high):
    store = SQLiteRunStore(tmp_path)
    runner = ValuationRunner(store, FinanceTeamModel())
    run = runner.create_run(relative_request(method))
    result = runner.execute(run.run_id)
    assert str(result.status).startswith("completed"), result.model_dump(mode="json")
    multiple = result.result.relative[0]
    assert (multiple.per_share_value, multiple.range_low, multiple.range_high) == (D(base), D(low), D(high))
    assert result.result.dcf is None and result.result.forecast == []
    package = build_valuation_bundle(store, result)
    replay = replay_bundle(package)
    assert replay["passed"], replay
    tampered = copy.deepcopy(package)
    tampered["effective_request"]["financials"]["common_shares"] = "1"
    with pytest.raises(ValueError, match="哈希"):
        replay_bundle(tampered)
    book = load_workbook(BytesIO(ValuationReportExporter().xlsx(result)))
    assert "预测与FCFF" not in book.sheetnames
    pdf = PdfReader(BytesIO(ValuationReportExporter().pdf(result)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert '预测与FCFF' not in text and 'WACC / g' not in text
    assert '敏感性项目与结果' in text


def test_dcf_missing_fields_is_blocked_not_zero_filled():
    request = relative_request().model_copy(update={"methods": ["dcf"]})
    findings = FinanceTeamModel().validate(request, request.financials)
    assert any(f.rule_id == "METHOD_INPUTS_MISSING" and f.severity == "blocking" for f in findings)


def test_research_pe_snapshot_does_not_require_four_years_or_fcff_fields():
    session = ResearchSession(session_id="research_relative", draft=draft(), data_source_preference="web")
    for metric, value, unit in [("net_income_parent", "200", "元"), ("common_shares", "100", "股")]:
        session.facts.append(FactCandidate(fact_id=metric, metric=metric, raw_value=value, normalized_value=value,
            unit=unit, period="2025", scope="consolidated", status="confirmed", block_id="file:1", quote=f"{metric} {value}"))
    assembler = ResearchValuationAssembler()
    assert assembler.structured_readiness_error(session) is None
    with pytest.raises(ValueError, match="可比样本"):
        assembler.build(session)
    for i in range(3):
        session.facts.append(FactCandidate(fact_id=f"peer_{i}", metric="pe", raw_value="10", normalized_value="10", unit="ratio",
            period="2026-06-30", scope="consolidated", status="confirmed", role="comparable", peer_name=f"同业{i}", peer_ticker=f"P{i}",
            multiple_basis="FY", block_id="file:1", quote="市盈率10倍"))
    request = assembler.build(session)
    assert request.financials.revenue is None and not request.historical_financials


def test_peer_evidence_requires_issuer_asof_and_matching_basis():
    item = candidate(role="comparable", metric="pe", unit="ratio", peer_name="测试同业", peer_ticker="600555", period="2026-06-30", multiple_basis="FY", quote="市盈率100倍")
    block = source("测试同业 600555 2026-06-30 年度口径 FY 市盈率100倍")
    assert not bind_evidence(item, block, [block], draft())[0]
    assert bind_evidence(item.model_copy(update={"multiple_basis": "TTM"}), block, [block], draft())[0]
    assert bind_evidence(item.model_copy(update={"period": "2025-06-30"}), block, [block], draft())[0]


def test_transport_timeout_retries_once_then_succeeds(monkeypatch):
    calls = []
    def handler(request):
        calls.append(request)
        if len(calls) == 1:
            raise httpx.ReadTimeout("simulated", request=request)
        return httpx.Response(200, json={"choices": [{"message": {"role": "assistant", "content": "ok"}}]})
    config = ModelConnectionInput(provider="openai_compatible", base_url="https://api.example.com/v1", model="test", api_key="test-only")
    http_client = httpx.Client
    monkeypatch.setattr(httpx, "Client", lambda **kwargs: http_client(transport=httpx.MockTransport(handler), **kwargs))
    client = OpenAICompatibleClient(config)
    assert client.chat([{"role": "user", "content": "hi"}])["content"] == "ok"
    assert len(calls) == 2


def test_dcf_bundle_replays_without_network_and_matches_independent_present_value(tmp_path, monkeypatch):
    store = SQLiteRunStore(tmp_path)
    runner = ValuationRunner(store, FinanceTeamModel())
    request = ValuationRequest(company={"name": "Synthetic independent DCF"}, valuation_date="2026-06-30", mode="demo", methods=["dcf"], discount_policy="year_end")
    record = runner.run(request)
    assert record.result is not None
    result = record.result
    wacc, g = result.assumptions.wacc, result.assumptions.terminal_growth
    pv = sum((row.fcff / (1 + wacc) ** row.discount_period for row in result.forecast), D(0))
    last = result.forecast[-1]
    terminal = last.fcff * (1 + g) / (wacc - g) / (1 + wacc) ** last.discount_period
    fin = result.effective_financials
    expected = (pv + terminal + fin.cash_and_non_operating_assets - fin.interest_bearing_debt) / fin.common_shares
    assert abs(expected - result.dcf.per_share_value) < D("0.0002")
    def no_network(*args, **kwargs):
        raise AssertionError("Offline replay attempted network I/O")
    monkeypatch.setattr(httpx.Client, "request", no_network)
    assert replay_bundle(build_valuation_bundle(store, record))["passed"]


def test_api_turn_survives_reload_and_duplicate_request_does_not_add_messages(tmp_path):
    app = create_app(tmp_path)
    with TestClient(app) as client:
        sid = client.post('/api/research-sessions', json={}).json()['session']['session_id']
        body = {"request_id": "request_api_once", "content": "说明估值方法"}
        assert client.post(f'/api/research-sessions/{sid}/turns', json=body).status_code == 202
        first = client.get(f'/api/research-sessions/{sid}?compact=true').json()
        assert first['execution']['status'] == 'completed'
        assert client.post(f'/api/research-sessions/{sid}/turns', json=body).status_code == 202
        second = client.get(f'/api/research-sessions/{sid}?compact=true').json()
        assert first['messages'] == second['messages']
        assert client.get(f'/api/research-sessions/{sid}').headers['cache-control'] == 'no-store'


def test_comparable_alias_normalization_is_explicit():
    assert candidate(role="comparable", metric="市盈率").metric == "pe"
    assert candidate(role="historical", metric="市盈率").metric == "市盈率"


def test_percent_comparison_column_does_not_shift_money_years_or_unit():
    block = source("测试股份 600123 合并利润表\n项目 2025年 2024年 本年比上年 2023年\n营业收入（元） 100 90 11.11% 80")
    warnings, verification = bind_evidence(candidate(unit="元", raw_value="80", period="2023"), block, [block], draft())
    assert not warnings and verification['year_column'] == 2023
    assert bind_evidence(candidate(unit="元", raw_value="80", period="2025"), block, [block], draft())[0]


def test_thousands_unit_is_not_guessed_as_yuan():
    block = source("测试股份 600123 2025年合并利润表\n单位：人民币千元\n营业收入 100")
    assert not bind_evidence(candidate(unit="千元"), block, [block], draft())[0]
    assert any('单位冲突' in w for w in bind_evidence(candidate(unit="元"), block, [block], draft())[0])


def test_multiple_years_in_one_prose_line_bind_amount_to_nearest_year():
    block = source("测试股份 600123 合并报表\n单位：万元\n2025年营业收入100万元，2024年营业收入90万元")
    assert not bind_evidence(candidate(), block, [block], draft())[0]
    assert any('年度行冲突' in w for w in bind_evidence(candidate(raw_value="90"), block, [block], draft())[0])


def test_a_number_elsewhere_in_the_block_is_not_evidence_for_this_metric():
    block = source("测试股份 600123 2025年合并报表\n单位：万元\n营业收入100\n归母净利润20")
    warnings, _ = bind_evidence(candidate(metric="归母净利润", raw_value="100"), block, [block], draft())
    assert any('科目数值冲突' in w for w in warnings)


def test_peer_multiple_cannot_belong_to_another_company_in_the_same_block():
    block = source("2026-06-30 年度FY 静态市盈率，单位：倍\n同业甲 600555 市盈率100倍\n同业乙 600666 市盈率20倍")
    item = candidate(role="comparable", metric="pe", unit="ratio", peer_name="同业乙", peer_ticker="600666", period="2026-06-30", multiple_basis="FY")
    assert any('可比数值未绑定' in w for w in bind_evidence(item, block, [block], draft())[0])


def test_xlsx_reference_recalc_preserves_stub_fraction_and_tax(tmp_path):
    runner = ValuationRunner(SQLiteRunStore(tmp_path), FinanceTeamModel())
    request = ValuationRequest(company={"name": "Synthetic DCF workbook"}, valuation_date="2026-06-30", mode="demo", methods=["dcf"])
    record = runner.run(request)
    book = load_workbook(BytesIO(ValuationReportExporter().xlsx(record)))
    sheet = book['DCF复算']
    assert D(str(sheet['F12'].value)) == record.result.effective_financials.tax_rate
    assert abs(D(str(sheet['O12'].value)) - record.result.forecast[0].cash_flow_fraction) < D('0.0000000001')
    assert sheet['N12'].value == '=K12*O12*M12'
    assert '/2)' in sheet['Q5'].value
    assert '*\'DCF复算\'!$O$12' in book['敏感性分析']['B2'].value


def test_xlsx_formal_recalc_excludes_operating_cash_and_treats_names_as_text(tmp_path):
    from test_finance_team_model import request
    req = request()
    req.company.name = '=HYPERLINK("https://invalid.example", "text only")'
    req.assumptions.operating_cash_ratio = D('0.02')
    record = ValuationRunner(SQLiteRunStore(tmp_path), FinanceTeamModel()).run(req)
    book = load_workbook(BytesIO(ValuationReportExporter().xlsx(record)))
    sheet = book['DCF复算']
    assert sheet['B9'].value > 0
    assert sheet['Q7'].value == '=MAX(0,$B$4-$B$9)'
    assert "MAX(0,'DCF复算'!$B$4-'DCF复算'!$B$9)" in book['敏感性分析']['B2'].value
    assert book['估值摘要']['B2'].data_type == 's'


def test_cancellation_and_resume_preserve_pending_fact_confirmation(tmp_path):
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create()
    fact = FactCandidate(fact_id='fact_pending', metric='revenue', raw_value='100', normalized_value='100', unit='元', period='2025', scope='consolidated', block_id='file:1', quote='营业收入100元')
    session.facts.append(fact)
    service._question(session, 'facts', '确认字段', [('accept', '确认'), ('defer', '稍后')], fact_ids=[fact.fact_id])
    service.store.save_research(session)
    turn = ResearchTurn(content='说明资料情况', request_id='cancel_pending_review')
    service.reserve_turn(session.session_id, turn)
    service.cancel_turn(session.session_id)
    stopped = service.turn(session.session_id, turn, reserved=True)
    assert stopped['session']['question']['kind'] == 'recovery'
    resumed = service.turn(session.session_id, ResearchTurn(question_id=stopped['session']['question']['question_id'], option_id='retry'))
    question = resumed['session']['question']
    assert question['kind'] == 'facts' and question['fact_ids'] == ['fact_pending']
    assert any(choice['id'] == 'accept' for choice in question['options'])
