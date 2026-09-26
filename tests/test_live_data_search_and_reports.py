import json
from datetime import date, datetime, timezone
from decimal import Decimal as D
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from valuationagent.api.main import create_app
from valuationagent.application.reporting import ValuationReportExporter
from valuationagent.application.research import ResearchService
from valuationagent.application.runner import ValuationRunner
from valuationagent.finance.team_model import FinanceTeamModel
from valuationagent.market.tushare import (
    TushareApiClient,
    TushareDataProvider,
    TusharePermissionError,
    map_tushare_industry,
    normalize_a_share_ticker,
)
from valuationagent.schemas.agent import SearchQuery
from valuationagent.schemas.models import (
    CompanyInput,
    FinancialSnapshot,
    RunStatus,
    ValuationRequest,
)
from valuationagent.schemas.research import FactCandidate
from valuationagent.search.providers import CninfoAnnouncementProvider, TavilySearchProvider
from valuationagent.storage.sqlite import SQLiteRunStore


def _request():
    history = []
    for year, revenue in enumerate(range(20, 30), 2016):
        amount = D(revenue) * D("100000000")
        history.append(FinancialSnapshot(
            period_end=date(year, 12, 31), revenue=amount, ebit_margin="0.20",
            tax_rate="0.15", depreciation_amortization=amount * D("0.04"),
            capital_expenditure=amount * D("0.05"), change_operating_nwc=amount * D("0.01"),
            cash_and_non_operating_assets="1000000000", interest_bearing_debt="200000000",
            common_shares="400000000", net_income_parent=amount * D("0.16"),
            ebitda=amount * D("0.24"),
            statement_items={"market_cap": "50000000000", "ebit": amount * D("0.20")},
            calculation_methods={
                "ebit_margin": "ebit / revenue",
                "ebitda": "ebit + depreciation_amortization",
                "ebit": "revenue * ebit_margin",
            },
        ))
    return ValuationRequest(
        company=CompanyInput(ticker="603893.SH", name="测试公司", industry="电子"),
        valuation_date=date(2026, 9, 23), forecast_years=10, discount_policy="year_end",
        methods=["dcf"], financials=history[-1], historical_financials=history[:-1],
    )


def test_tavily_provider_uses_cutoff_and_preserves_url_without_leaking_key():
    captured = {}

    def handler(request):
        captured["authorization"] = request.headers["Authorization"]
        captured["payload"] = __import__("json").loads(request.content)
        return httpx.Response(200, json={"results": [{
            "title": "Company annual report", "url": "https://example.com/report",
            "content": "Audited annual figures", "score": 0.9, "published_date": "2025-04-01",
        }]})

    provider = TavilySearchProvider("secret-test-key", transport=httpx.MockTransport(handler))
    result = provider.search(SearchQuery(
        query="测试公司 年报", purpose="financials", information_cutoff=date(2025, 12, 31)
    ))
    assert result.status == "completed"
    assert result.hits[0].url == "https://example.com/report"
    assert captured["payload"]["end_date"] == "2025-12-31"
    assert captured["authorization"] == "Bearer secret-test-key"
    assert "secret-test-key" not in result.model_dump_json()


def test_tavily_auth_failure_is_distinct_and_never_leaks_key():
    def handler(request):
        return httpx.Response(401, json={"detail": "invalid credential"})

    secret = "secret-test-key"
    provider = TavilySearchProvider(secret, transport=httpx.MockTransport(handler))
    result = provider.search(SearchQuery(query="测试公司 年报", purpose="financials"))
    assert result.status == "failed"
    assert result.error_code == "SEARCH_AUTH_FAILED"
    assert "重新配置" in result.error_message
    assert secret not in result.model_dump_json()


def test_tavily_transient_failure_is_retried_without_user_intervention():
    attempts = 0

    def handler(request):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            return httpx.Response(503, json={"detail": "temporary"})
        return httpx.Response(200, json={"results": [{
            "title": "Policy source",
            "url": "https://example.com/policy",
            "content": "Official policy text",
            "score": 0.8,
        }]})

    result = TavilySearchProvider(
        "secret-test-key", transport=httpx.MockTransport(handler)
    ).search(SearchQuery(query="industry policy", purpose="policy"))

    assert result.status == "completed"
    assert attempts == 2
    assert any("自动重试 1 次后成功" in warning for warning in result.warnings)


def test_cninfo_catalogue_returns_exact_full_annual_reports_by_code_and_year():
    captured = {}

    def announcement(identifier, code, title, path, published):
        return {
            "secCode": code,
            "secName": "<em>恒瑞医药</em>",
            "announcementId": identifier,
            "announcementTitle": title,
            "announcementTime": int(
                datetime.combine(published, datetime.min.time(), tzinfo=timezone.utc).timestamp() * 1000
            ),
            "adjunctUrl": path,
        }

    def handler(request):
        from urllib.parse import parse_qs
        captured.update({key: value[0] for key, value in parse_qs(request.content.decode()).items()})
        return httpx.Response(200, json={"announcements": [
            announcement("full24", "600276", "<em>恒瑞医药</em>2024年年度报告", "finalpage/2025-03-31/full24.PDF", date(2025, 3, 31)),
            announcement("summary24", "600276", "恒瑞医药2024年年度报告摘要", "finalpage/2025-03-31/summary24.PDF", date(2025, 3, 31)),
            announcement("full23", "600276", "恒瑞医药2023年年度报告", "finalpage/2024-04-18/full23.PDF", date(2024, 4, 18)),
            announcement("other", "600519", "贵州茅台2024年年度报告", "finalpage/2025-03-31/other.PDF", date(2025, 3, 31)),
        ]})

    provider = CninfoAnnouncementProvider(transport=httpx.MockTransport(handler))
    result = provider.search_annual_reports(
        "600276.SH", [2023, 2024], cutoff=date(2025, 4, 1), company_name="恒瑞医药"
    )

    assert result.status == "completed"
    assert [hit.source_id for hit in result.hits] == ["cninfo_full24", "cninfo_full23"]
    assert result.hits[0].published_at == date(2025, 3, 31)
    assert all(hit.domain == "static.cninfo.com.cn" for hit in result.hits)
    assert all("摘要" not in hit.title for hit in result.hits)
    assert captured["searchkey"] == "600276"
    assert captured["category"] == "category_ndbg_szsh;"
    assert captured["seDate"] == "2024-01-01~2025-04-01"


def test_cninfo_catalogue_retries_exact_issuer_for_shenzhen_main_board():
    from urllib.parse import parse_qs

    requests = []

    def handler(request):
        data = {key: value[0] for key, value in parse_qs(request.content.decode()).items()}
        requests.append(data)
        if data.get("stock") != "000858,gssz0000858":
            return httpx.Response(200, json={"announcements": []})
        return httpx.Response(200, json={"announcements": [{
            "secCode": "000858",
            "secName": "五粮液",
            "announcementId": "wly2024",
            "announcementTitle": "五粮液2024年年度报告",
            "announcementTime": int(datetime(2025, 4, 26, tzinfo=timezone.utc).timestamp() * 1000),
            "adjunctUrl": "finalpage/2025-04-26/wly2024.PDF",
        }]})

    result = CninfoAnnouncementProvider(
        transport=httpx.MockTransport(handler)
    ).search_annual_reports("000858.SZ", [2024], cutoff=date(2026, 9, 24))

    assert result.status == "completed"
    assert result.hits[0].source_id == "cninfo_wly2024"
    assert requests[0]["searchkey"] == "000858"
    assert requests[1]["stock"] == "000858,gssz0000858"


def test_tushare_client_and_industry_mapping_are_deterministic():
    def handler(request):
        body = __import__("json").loads(request.content)
        assert body["token"] == "token-value"
        return httpx.Response(200, json={"code": 0, "data": {
            "fields": ["ts_code", "name"], "items": [["600519.SH", "贵州茅台"]]
        }})

    client = TushareApiClient("token-value", transport=httpx.MockTransport(handler))
    assert client.query("stock_basic")[0]["name"] == "贵州茅台"
    assert normalize_a_share_ticker("600519") == "600519.SH"
    label, warning = map_tushare_industry("白酒")
    assert label == "轻工制造"
    assert "显式映射" in warning


def test_tushare_permission_failure_is_typed_without_exposing_provider_message():
    def handler(request):
        return httpx.Response(200, json={"code": -2001, "msg": "抱歉，您没有接口访问权限"})

    client = TushareApiClient("token-value", transport=httpx.MockTransport(handler))
    with pytest.raises(TusharePermissionError, match="stock_basic.*接口权限不足"):
        client.query("stock_basic")


def test_stock_basic_permission_can_fall_back_to_confirmed_company_metadata(tmp_path, monkeypatch):
    class PermissionLimitedClient:
        def query(self, api_name, *, params=None, fields=None):
            if api_name == "stock_basic":
                raise TusharePermissionError(api_name, "接口权限不足")
            raise AssertionError(api_name)

    request = _request().model_copy(update={"data_source": "ticker"})
    provider = TushareDataProvider(PermissionLimitedClient())
    history = [*request.historical_financials, request.financials]
    monkeypatch.setattr(provider, "_latest_daily", lambda ticker, cutoff: {
        "trade_date": "20260923", "total_share": 40000, "total_mv": 5000000,
    })
    monkeypatch.setattr(provider, "_market_cap_statistics", lambda ticker, cutoff: {})
    monkeypatch.setattr(provider, "_financial_snapshots", lambda *args: (history, []))

    bundle = provider.resolve(request, SQLiteRunStore(tmp_path / "fallback"))

    assert bundle.company.name == "测试公司"
    assert bundle.company.ticker == "603893.SH"
    assert bundle.company.exchange == "SSE"
    assert any("采用研究会话中已确认的信息" in warning for warning in bundle.warnings)


def test_stock_basic_permission_requires_confirmed_industry_before_fcff(tmp_path):
    class PermissionLimitedClient:
        def query(self, api_name, *, params=None, fields=None):
            raise TusharePermissionError(api_name, "接口权限不足")

    request = _request().model_copy(update={
        "data_source": "ticker",
        "company": CompanyInput(ticker="603893.SH", name="测试公司"),
    })
    provider = TushareDataProvider(PermissionLimitedClient())

    with pytest.raises(ValueError, match="确认该公司的非金融行业"):
        provider.resolve(request, SQLiteRunStore(tmp_path / "missing-industry"))


def test_peer_permission_failure_keeps_dcf_and_marks_relative_methods_unavailable(tmp_path, monkeypatch):
    class PermissionLimitedClient:
        def query(self, api_name, *, params=None, fields=None):
            if api_name == "stock_basic":
                raise TusharePermissionError(api_name, "接口权限不足")
            raise AssertionError(api_name)

    request = _request().model_copy(update={
        "data_source": "ticker", "methods": ["dcf", "pe", "ev_ebitda"],
    })
    provider = TushareDataProvider(PermissionLimitedClient())
    history = [*request.historical_financials, request.financials]
    monkeypatch.setattr(provider, "_latest_daily", lambda ticker, cutoff: {
        "trade_date": "20260923", "total_share": 40000, "total_mv": 5000000,
    })
    monkeypatch.setattr(provider, "_market_cap_statistics", lambda ticker, cutoff: {})
    monkeypatch.setattr(provider, "_financial_snapshots", lambda *args: (history, []))

    bundle = provider.resolve(request, SQLiteRunStore(tmp_path / "peer-fallback"))

    assert bundle.peers == []
    assert any("DCF继续执行" in warning and "样本不足" in warning for warning in bundle.warnings)


def test_peer_selection_excludes_financial_and_uses_cash_in_enterprise_value():
    class FakeClient:
        def query(self, api_name, *, params=None, fields=None):
            if api_name == "stock_basic":
                return [
                    {"ts_code": "A.SZ", "name": "同行A", "industry": "软件服务"},
                    {"ts_code": "B.SZ", "name": "同行B", "industry": "软件服务"},
                    {"ts_code": "C.SZ", "name": "同行C", "industry": "软件服务"},
                    {"ts_code": "F.SH", "name": "银行F", "industry": "银行"},
                ]
            if api_name == "daily_basic":
                return [{"ts_code": code, "trade_date": "20260923", "pe_ttm": 20,
                         "ps_ttm": 4, "total_mv": 1000000}
                        for code in ("A.SZ", "B.SZ", "C.SZ")]
            if api_name == "fina_indicator":
                return [{"ts_code": params["ts_code"], "ann_date": "20260401",
                         "end_date": "20251231", "or_yoy": 10, "ebit_of_gr": 20,
                         "ebitda": 1000000000, "interestdebt": 200000000}]
            if api_name == "balancesheet":
                return [{"ts_code": params["ts_code"], "ann_date": "20260401",
                         "end_date": "20251231", "money_cap": 100000000}]
            raise AssertionError(api_name)

    provider = TushareDataProvider(FakeClient(), peer_limit=3)
    peers, warnings = provider._select_peers(
        CompanyInput(ticker="TARGET.SZ", industry="软件服务"),
        {"trade_date": "20260923", "total_mv": 1000000}, date(2026, 9, 23),
    )
    assert not warnings
    assert [peer.ticker for peer in peers] == ["A.SZ", "B.SZ", "C.SZ"]
    # total_mv uses ten-thousand yuan; EV=(10bn+0.2bn-0.1bn)/1bn=10.1x.
    assert peers[0].ev_ebitda == D("10.1")
    assert peers[0].market_cap == D("10000000000")
    assert peers[0].revenue_growth == D("0.1")
    assert peers[0].ebit_margin == D("0.2")
    assert peers[0].peer_tier == "core"


def test_market_cap_statistics_preserve_average_and_period_endpoints():
    class FakeClient:
        def query(self, api_name, *, params=None, fields=None):
            assert api_name == "daily_basic"
            rows = []
            for prefix, total_mv in (
                ("202510", 100),
                ("202601", 200),
                ("202604", 300),
                ("202607", 400),
            ):
                rows.extend({
                    "ts_code": "TARGET.SZ",
                    "trade_date": f"{prefix}{day:02d}",
                    "total_mv": total_mv,
                } for day in range(1, 21))
            return rows

    statistics = TushareDataProvider(FakeClient())._market_cap_statistics(
        "TARGET.SZ", date(2026, 9, 23)
    )
    assert statistics["quarterly_average_market_cap"] == D("2500000")
    assert statistics["annual_average_market_cap"] == D("2500000")
    assert statistics["market_cap_period_low"] == D("1000000")
    assert statistics["market_cap_period_high"] == D("4000000")


def test_research_ticker_handoff_creates_traceable_run(tmp_path):
    store = SQLiteRunStore(tmp_path / "handoff")
    research = ResearchService(store)
    runner = ValuationRunner(store, FinanceTeamModel())
    session = research.create()
    session.data_source_preference = "online"
    session.draft = session.draft.model_copy(update={
        "company": "贵州茅台", "ticker": "600519.SH", "valuation_date": date(2026, 9, 23),
        "methods": ["dcf", "pe", "ps", "ev_ebitda"],
    })
    store.save_research(session)
    session_provider = TushareDataProvider(TushareApiClient("SESSION_ONLY_TOKEN"))
    research.attach_market(session.session_id, session_provider)
    record = research.submit_valuation(session.session_id, runner)
    saved = store.get_research(session.session_id)
    assert record.request.data_source == "ticker"
    assert record.request.forecast_years == 10
    assert saved.valuation_run_id == record.run_id
    assert runner._run_data[record.run_id] is session_provider
    assert any(event.type == "valuation.submitted" for event in store.list_events(session.session_id))


def test_research_handoff_creates_a_new_run_after_confirmed_inputs_change(tmp_path):
    store = SQLiteRunStore(tmp_path / "handoff-revision")
    research = ResearchService(store)
    runner = ValuationRunner(store, FinanceTeamModel())
    session = research.create()
    session.data_source_preference = "online"
    session.draft = session.draft.model_copy(update={
        "company": "测试公司", "ticker": "603893.SH",
        "valuation_date": date(2026, 9, 23), "methods": ["dcf"],
    })
    store.save_research(session)
    first = research.submit_valuation(session.session_id, runner)
    store.update_run(first.run_id, status=RunStatus.FAILED, error={"message": "metadata unavailable"})

    updated = store.get_research(session.session_id)
    updated.draft.industry = "电子 / 计算机 / 半导体"
    store.save_research(updated)
    second = research.submit_valuation(session.session_id, runner)

    assert second.run_id != first.run_id
    assert second.parent_run_id == first.run_id
    assert second.request.company.industry == "电子 / 计算机 / 半导体"
    assert second.revision_reason == "研究会话输入更新后重新提交"


def test_web_data_services_are_session_scoped_and_secrets_are_not_persisted(tmp_path):
    app = create_app(tmp_path / "data-services")
    session = app.state.research.create()
    tavily_key = "tvly-SESSION_ONLY_TEST_KEY"
    tushare_token = "TUSHARE_SESSION_ONLY_TEST_TOKEN"
    with TestClient(app) as client:
        before = client.get(f"/api/research-sessions/{session.session_id}/data-services")
        connected = client.post(
            f"/api/research-sessions/{session.session_id}/data-services",
            json={"tavily_api_key": tavily_key, "tushare_token": tushare_token},
        )
        after = client.get(f"/api/research-sessions/{session.session_id}/data-services")
    assert before.status_code == 200
    assert before.json()["search"]["available"] is False
    assert before.json()["market"]["available"] is False
    assert connected.status_code == 200
    assert after.json()["search"] == {"available": True, "provider": "tavily", "connection_status": "configured"}
    assert after.json()["market"]["available"] is True
    persisted = json.dumps(app.state.research.snapshot(session.session_id), ensure_ascii=False)
    assert tavily_key not in persisted
    assert tushare_token not in persisted


def test_upload_preference_never_silently_falls_back_to_ticker(tmp_path):
    store = SQLiteRunStore(tmp_path / "upload-preference")
    research = ResearchService(store)
    runner = ValuationRunner(store, FinanceTeamModel())
    session = research.create()
    session.data_source_preference = "upload"
    session.draft = session.draft.model_copy(update={
        "company": "贵州茅台",
        "ticker": "600519.SH",
        "valuation_date": date(2026, 9, 23),
        "methods": ["dcf"],
    })
    store.save_research(session)
    with pytest.raises(ValueError, match="已选择自行上传"):
        research.submit_valuation(session.session_id, runner)


def test_formal_xlsx_and_pdf_reports_are_readable(tmp_path):
    store = SQLiteRunStore(tmp_path / "reports")
    record = ValuationRunner(store, FinanceTeamModel()).run(_request())
    assert record.result is not None
    exporter = ValuationReportExporter()
    xlsx = exporter.xlsx(record)
    workbook = load_workbook(BytesIO(xlsx), data_only=False)
    assert {"估值摘要", "关键假设", "基期推导", "预测与FCFF", "DCF复算", "敏感性分析", "来源与风险"} <= set(workbook.sheetnames)
    assert "可比公司" not in workbook.sheetnames
    assert str(workbook["预测与FCFF"]["I2"].value).startswith("=")
    assert str(workbook["DCF复算"]["Q10"].value).startswith("=")
    assert str(workbook["敏感性分析"]["B2"].value).startswith("=IF(")
    derivation_rows = list(workbook["基期推导"].iter_rows(values_only=True))
    assert any(row[1] == "EBITDA (ebitda)" and row[3] == "ebit + depreciation_amortization" for row in derivation_rows)
    assert any(row[1] == "ebit" and row[3] == "revenue * ebit_margin" for row in derivation_rows)
    assert all(sheet.max_row > 1 for sheet in workbook.worksheets)
    pdf = exporter.pdf(record)
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) >= 3


def test_confirmed_raw_statement_lines_complete_valuation_and_auditable_report(tmp_path):
    store = SQLiteRunStore(tmp_path / "derived-e2e")
    research = ResearchService(store)
    session = research.create()
    session.draft = session.draft.model_copy(update={
        "company": "测试汽车",
        "industry": "汽车制造",
        "valuation_date": date(2025, 4, 30),
        "methods": ["dcf"],
    })
    session.data_source_preference = "web"
    revenues = {
        2021: D("750000000"),
        2022: D("820000000"),
        2023: D("900000000"),
        2024: D("1000000000"),
    }
    for year, revenue in revenues.items():
        profit_before_tax = revenue * D("0.10")
        inputs = {
            "revenue": revenue,
            "profit_before_tax": profit_before_tax,
            "income_tax_expense": profit_before_tax * D("0.15"),
            "interest_expense": revenue * D("0.01"),
            "depreciation_fixed_assets": revenue * D("0.03"),
            "amortization_intangible_assets": revenue * D("0.005"),
            "amortization_long_term_deferred_expenses": revenue * D("0.005"),
            "depreciation_right_of_use": D("0"),  # Explicit synthetic disclosure.
            "cash_paid_for_ppe_intangibles": revenue * D("0.05"),
            "inventory_decrease": -(revenue * D("0.003")),
            "operating_receivables_decrease": -(revenue * D("0.004")),
            "operating_payables_increase": -(revenue * D("0.003")),
            "cash_and_non_operating_assets": revenue * D("0.20"),
            "short_term_borrowings": revenue * D("0.04"),
            "current_portion_non_current_liabilities": revenue * D("0.01"),
            "long_term_borrowings": revenue * D("0.05"),
            "bonds_payable": revenue * D("0.02"),
            "lease_liabilities": revenue * D("0.01"),
            "common_shares": D("100000000"),
            "net_income_parent": revenue * D("0.085"),
        }
        for metric, value in inputs.items():
            session.facts.append(FactCandidate(
                fact_id=f"fact_{metric}_{year}",
                metric=metric,
                raw_value=str(value),
                unit="股" if metric == "common_shares" else "元",
                normalized_value=str(value),
                period=str(year),
                scope="consolidated",
                block_id=f"message:{metric}_{year}",
                quote=f"{year} {metric} {value}",
                status="confirmed",
            ))

    request = research.valuation_assembler.build(session)
    assert len(request.historical_financials) == 3
    assert request.financials.calculation_methods["ebitda"] == (
        "ebit + depreciation_amortization"
    )

    record = ValuationRunner(store, FinanceTeamModel()).run(request)

    assert str(record.status) == "completed_with_warnings"
    assert record.result is not None and record.result.dcf is not None
    assert len(record.result.forecast) == 10
    assert record.result.sensitivity
    workbook = load_workbook(
        BytesIO(ValuationReportExporter().xlsx(record)), data_only=False
    )
    rows = list(workbook["基期推导"].iter_rows(values_only=True))
    assert any(
        row[1] == "EBITDA (ebitda)"
        and row[3] == "ebit + depreciation_amortization"
        for row in rows
    )


def test_api_exposes_research_handoff_and_formal_exports(tmp_path):
    app = create_app(tmp_path / "api")
    record = app.state.runner.run(_request())
    research_session = app.state.research.create()
    research_session.data_source_preference = "online"
    research_session.draft = research_session.draft.model_copy(update={
        "company": "贵州茅台", "ticker": "600519.SH", "valuation_date": date(2026, 9, 23),
        "methods": ["dcf"],
    })
    app.state.store.save_research(research_session)
    with TestClient(app) as client:
        handoff = client.post(f"/api/research-sessions/{research_session.session_id}/valuation")
        assert handoff.status_code == 202
        xlsx = client.get(f"/api/runs/{record.run_id}/export?format=xlsx")
        pdf = client.get(f"/api/runs/{record.run_id}/export?format=pdf")
    assert xlsx.status_code == 200 and xlsx.content.startswith(b"PK")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
