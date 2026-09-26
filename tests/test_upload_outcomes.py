"""Real file parsers + evidence gate, not a mocked claim that upload succeeded."""
import csv
import json
from datetime import date
from io import BytesIO, StringIO

import pytest
from pydantic import ValidationError
from pypdf import PdfWriter

from valuationagent.application.research import CandidateInput, ProposeFacts, ResearchService
from valuationagent.application.research_export import build_research_export
from valuationagent.application.result_document import _report_font
from valuationagent.schemas.research import ResearchDraft, ResearchTurn
from valuationagent.search.providers import MockSearchProvider
from valuationagent.storage.sqlite import SQLiteRunStore


ROWS = [["测试股份 600123 合并利润表"], ["单位：万元"], ["项目", "2025年", "2024年"], ["营业收入", "100", "90"]]


def payload(extension):
    """Synthetic statement, same financial content in all advertised formats."""
    text = "\n".join(" | ".join(row) for row in ROWS)
    data = BytesIO()
    if extension == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        pdf = Canvas(data)
        pdf.setFont(_report_font(), 12)
        for i, line in enumerate(text.splitlines()):
            pdf.drawString(48, 780 - i * 24, line)
        pdf.save()
    elif extension == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph(ROWS[0][0])
        document.add_paragraph(ROWS[1][0])
        table = document.add_table(rows=2, cols=3)
        for r, row in enumerate(ROWS[2:]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        document.save(data)
    elif extension == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        for row in ROWS:
            book.active.append(row)
        book.save(data)
    elif extension in {"csv", "tsv"}:
        csv_data = StringIO(newline="")
        csv.writer(csv_data, delimiter="," if extension == "csv" else "\t").writerows(ROWS)
        return csv_data.getvalue().encode("utf-8-sig")
    elif extension == "html":
        return ("<html><body><script>fake_revenue=999</script><h1>" + ROWS[0][0] + "</h1><p>" + ROWS[1][0] +
                "</p><table>" + "".join("<tr>" + "".join("<td>" + cell + "</td>" for cell in row) + "</tr>" for row in ROWS[2:]) + "</table></body></html>").encode()
    elif extension == "json":
        return json.dumps({"title": ROWS[0][0], "unit": ROWS[1][0], "table": [" | ".join(row) for row in ROWS[2:]]}, ensure_ascii=False).encode()
    else:
        return text.encode()
    return data.getvalue()


@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx", "csv", "tsv", "html", "json", "txt", "md"])
def test_uploaded_formats_reach_evidence_checked_candidate(tmp_path, extension):
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create(data_source_preference="web")
    session.draft = ResearchDraft(company="测试股份", ticker="600123", industry="电子", methods=["ps"], valuation_date=date(2026, 6, 30))
    service.store.save_research(session)
    meta = service.store.save_upload("财务." + extension, "historical_financials", None, payload(extension))
    state = service.turn(session.session_id, ResearchTurn(content="读取附件", file_ids=[meta["file_id"]]))
    assert state["session"]["data_source_preference"] == "web"
    assert state["session"]["documents"][0]["parse_status"] == "parsed"
    blocks = service.store.research_blocks(session.session_id, meta["file_id"])
    row = next(block for block in blocks if "营业收入" in block["text"])
    assert row["location"] and "fake_revenue" not in row["text"]
    session = service.store.get_research(session.session_id)
    service._facts(session, ProposeFacts(candidates=[CandidateInput(metric="营业收入", raw_value="100", unit="万元", period="2025", scope="consolidated", block_id=row["block_id"], quote=row["text"])]))
    assert len(session.facts) == 1
    fact = session.facts[0]
    assert not fact.warnings, fact.warnings
    assert fact.normalized_value == "1000000" and fact.source_sha256 == meta["sha256"]
    assert fact.verification["year_column"] == 2025
    assert fact.status == "proposed"  # Extraction is not authorization to calculate.


@pytest.mark.parametrize("kind", ["broken_pdf", "encrypted_pdf", "no_text_pdf", "bad_docx", "bad_xlsx", "bad_json"])
def test_bad_upload_is_isolated_other_files_and_report_survive(tmp_path, kind):
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create(data_source_preference="web")
    session.draft = ResearchDraft(company="测试股份", methods=["dcf"], valuation_date=date(2026, 6, 30))
    service.store.save_research(session)
    suffix = kind.rsplit("_", 1)[-1]
    bad = b"not a valid financial document"
    if kind in {"encrypted_pdf", "no_text_pdf"}:
        writer = PdfWriter()
        writer.add_blank_page(width=500, height=700)
        if kind == "encrypted_pdf":
            writer.encrypt("fixture-only")
        stream = BytesIO()
        writer.write(stream)
        bad = stream.getvalue()
    uploads = [service.store.save_upload("坏资料." + suffix, "historical_financials", None, bad),
               service.store.save_upload("可读资料.txt", "historical_financials", None, payload("txt"))]
    state = service.turn(session.session_id, ResearchTurn(content="开始估值", file_ids=[f["file_id"] for f in uploads]))
    assert [d["parse_status"] for d in state["session"]["documents"]] == ["unreadable", "parsed"]
    assert state["session"]["documents"][0]["warnings"]
    assert state["result_document"]["status"] == "insufficient_data"
    assert not state["result_document"]["numeric_result_available"]
    html, _ = build_research_export(service, session.session_id, "html")
    assert "坏资料" in html and "读取限制" in html


def test_twenty_one_candidates_are_validated_without_truncation(tmp_path):
    from test_delivery_regressions import candidate, source
    from valuationagent.schemas.research import DocumentSummary
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create()
    candidates = []
    blocks = []
    for i in range(21):
        year = 2000 + i
        block = source(f"测试股份 合并报表\n单位：万元\n项目 | {year}年 | {year - 1}年\n营业收入 | 100 | 90", page=i + 1)
        block["block_id"] = f"file_1:{i + 1}"
        blocks.append(block)
        candidates.append(candidate(period=str(year), block_id=block["block_id"]))
    service.store.save_research_blocks(session.session_id, "file_1", blocks)
    session.documents.append(DocumentSummary(file_id="file_1", name="batch.txt", role="historical_financials", block_count=21))
    session.draft.company = "测试股份"
    service._facts(session, ProposeFacts(candidates=candidates))
    assert len(session.facts) == 21
    assert all(not fact.warnings for fact in session.facts)
    with pytest.raises(ValidationError):
        ProposeFacts(candidates=candidates * 4)


def test_failed_attachment_automatically_uses_allowed_public_retrieval(tmp_path):
    from test_evidence_recovery import RepeatingRetrievalModel
    model = RepeatingRetrievalModel()
    provider = MockSearchProvider()
    service = ResearchService(SQLiteRunStore(tmp_path), search_provider=provider)
    session = service.create(llm=model, data_source_preference="web")
    session.draft = ResearchDraft(company="测试股份", industry="电子", methods=["dcf"], valuation_date=date(2026, 6, 30))
    service.store.save_research(session)
    meta = service.store.save_upload("不可读.pdf", "historical_financials", None, b"not pdf")
    state = service.turn(session.session_id, ResearchTurn(content="开始估值", file_ids=[meta["file_id"]]))
    assert len(state["session"]["search_history"]) == 3
    assert state["result_document"] and model.calls == 4
    assert state["session"]["question"]["options"][0]["id"] == "report"


def test_explicit_upload_only_still_prohibits_search(tmp_path):
    from test_evidence_recovery import RepeatingRetrievalModel
    service = ResearchService(SQLiteRunStore(tmp_path), search_provider=MockSearchProvider())
    session = service.create(llm=RepeatingRetrievalModel(), data_source_preference="upload")
    session.draft = ResearchDraft(company="测试股份", industry="电子", methods=["dcf"], valuation_date=date(2026, 6, 30))
    service.store.save_research(session)
    state = service.turn(session.session_id, ResearchTurn(content="开始估值"))
    assert not state["session"]["search_history"]
    assert state["result_document"]


@pytest.mark.parametrize("explicit_target", [False, True])
def test_peer_annual_search_does_not_reuse_target_company_code(tmp_path, explicit_target):
    from test_research_sessions import ScriptedModel
    from valuationagent.schemas.agent import SearchQuery, SearchResult
    class Catalogue:
        def __init__(self):
            self.seen = []
        def search_annual_reports(self, ticker, years, *, cutoff, company_name):
            self.seen.append((ticker, company_name))
            return SearchResult(query=SearchQuery(query="annual report", ticker=ticker), provider="catalogue", provider_version="test", status="no_results")
    catalogue = Catalogue()
    args = {"query": "五粮液2024年度报告" if explicit_target else "000858五粮液2024年度报告", "purpose": "financials", "reason": "查找可比公司年报"}
    if explicit_target:
        args["target_ticker"] = "000858.SZ"
    model = ScriptedModel([("search_sources", args), ("finish_response", {"answer": "未取得相应原文，不填入数字。"})])
    service = ResearchService(SQLiteRunStore(tmp_path), search_provider=MockSearchProvider(), official_search_provider=catalogue)
    session = service.create(llm=model, data_source_preference="web")
    session.draft = ResearchDraft(company="贵州茅台", ticker="600519.SH", valuation_date=date(2025, 6, 30))
    service.store.save_research(session)
    service.turn(session.session_id, ResearchTurn(content="查找同业年报"))
    assert catalogue.seen == [("000858.SZ" if explicit_target else "000858", None)]


def test_keyword_retrieval_ranks_statement_rows_and_returns_headers(tmp_path):
    from test_research_sessions import ScriptedModel
    from valuationagent.schemas.research import DocumentSummary
    model = ScriptedModel([("read_document", {"file_id": "file_test", "query": "固定资产折旧 无形资产摊销", "limit": 1}),
                           ("finish_response", {"answer": "已读取科目及其原文表头。"})])
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create(llm=model)
    blocks = [{"block_id": f"file_test:{i + 1}", "file_id": "file_test", "text": text, "location": {"page": i + 1}}
              for i, text in enumerate(["目录 固定资产折旧与无形资产摊销 120", "合并现金流量表补充资料\n单位：万元\n项目 | 2025年 | 2024年", "固定资产折旧 | 100.00 | 90.00\n无形资产摊销 | 20.00 | 18.00"])]
    service.store.save_research_blocks(session.session_id, "file_test", blocks)
    session.documents.append(DocumentSummary(file_id="file_test", name="annual.pdf", role="historical_financials", block_count=3))
    service.store.save_research(session)
    state = service.turn(session.session_id, ResearchTurn(content="查看折旧摊销原文"))
    result = next(event["payload"]["output"] for event in state["events"] if event["type"] == "tool.completed" and event["tool"] == "read_document")
    assert result["blocks"][0]["block_id"] == "file_test:3"
    assert any(block["block_id"] == "file_test:2" for block in result["context_blocks"])


def test_unmatched_correction_preserves_old_fact_but_not_block_other_valid_candidates(tmp_path):
    from test_delivery_regressions import candidate, source
    from valuationagent.schemas.research import FactCandidate, DocumentSummary
    service = ResearchService(SQLiteRunStore(tmp_path))
    session = service.create()
    block = source("测试股份 合并报表\n单位：万元\n项目 | 2025年 | 2024年\n营业收入 | 100 | 90")
    service.store.save_research_blocks(session.session_id, "file_1", [block])
    session.documents.append(DocumentSummary(file_id="file_1", name="test.txt", role="historical_financials", block_count=1))
    old = FactCandidate(fact_id="old_confirmed", metric="所得税费用", period="2025", raw_value="20", unit="万元", normalized_value="200000", scope="consolidated", block_id="file_1:1", quote="old", status="confirmed")
    session.facts.append(old)
    service._facts(session, ProposeFacts(candidates=[candidate()], replaces=[old.fact_id, "unknown_id"]))
    assert len(session.facts) == 2 and not session.facts[-1].warnings
    assert old.status == "confirmed" and old.normalized_value == "200000"
    assert not session.question.superseded_fact_ids
    assert any(e.type == "facts.correction_rejected" for e in service.store.list_events(session.session_id))
