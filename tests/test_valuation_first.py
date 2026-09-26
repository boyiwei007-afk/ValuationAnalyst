"""Synthetic, offline acceptance of continuous modeling, not live-LLM accuracy."""
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from test_evidence_recovery import block, configured_service, item
from test_finance_team_model import history
from test_research_sessions import ScriptedModel

from valuationagent.application.reporting import ValuationReportExporter
from valuationagent.application.reproducibility import (
    build_valuation_bundle,
    replay_bundle,
)
from valuationagent.application.research import (
    ProposeFacts,
    ProposeForecast,
    _explicit_fact_acceptance,
)
from valuationagent.application.runner import ValuationRunner
from valuationagent.application.valuation_plan import valuation_progress
from valuationagent.finance.team_model import FinanceTeamModel
from valuationagent.llm.intent import requests_valuation
from valuationagent.schemas.models import required_financial_metrics
from valuationagent.schemas.research import (
    DocumentSummary,
    ForecastInputs,
    ResearchTurn,
)


def forecast(evidence_id="file_test:1"):
    return {"inputs": {"revenue_growth_scenarios": {
        "pessimistic": ["0.01"] * 10, "base": ["0.05"] * 10, "optimistic": ["0.08"] * 10},
        "ebit_margin_scenarios": {"pessimistic": ["0.18"] * 10, "base": ["0.24"] * 10, "optimistic": ["0.27"] * 10},
        "wacc": "0.10", "terminal_growth": "0.02"},
        "rationale": "这是合成测试的预测判断，不是历史事实；以已核验基期为锚，设置不同需求增长及利润率路径，并明确折现率与终值假设。",
        "risks": ["历史样本仅一年，无法验证增长趋势；预测及折现率存在主观性。"],
        "evidence_ids": [evidence_id]}


def baseline(service, session, *, omit=(), methods=("dcf",)):
    row = history()[-1]
    fields = sorted(required_financial_metrics(methods) - set(omit))
    units = {key: "ratio" if key in {"ebit_margin", "tax_rate"} else "股" if key == "common_shares" else "元" for key in fields}
    lines = {key: f"{key}（{units[key]}） {getattr(row, key)}" for key in fields}
    source = block(f"{session.draft.company} {session.draft.ticker} 2025年合并报表\n单位：元\n" + "\n".join(lines.values()))
    meta = service.store.save_upload("synthetic-baseline.txt", "historical_financials", "text/plain", source["text"].encode())
    fid = meta["file_id"]
    source.update({"block_id": fid + ":1", "file_id": fid})
    session.documents = [DocumentSummary(file_id=fid, name=meta["original_name"], role="historical_financials",
                                        block_count=1, sha256=meta["sha256"], size_bytes=meta["size_bytes"])]
    service.store.save_research_blocks(session.session_id, fid, [source])
    return [item(key, str(getattr(row, key)), block_id=fid + ":1", unit=units[key], quote=lines[key]) for key in fields]


@pytest.mark.parametrize("company,ticker", [("样本制造甲公司", "600123"), ("样本制造乙公司", "000321"),
                                           ("样本制造丙公司", "300123"), ("样本制造丁公司", "688123")])
def test_multiple_batches_to_one_review_to_dcf_sensitivity_reports_and_replay(tmp_path, company, ticker):
    service, session, _ = configured_service(tmp_path)
    session.draft.company, session.draft.ticker = company, ticker
    candidates = [f.model_dump() for f in baseline(service, session)]
    model = ScriptedModel([
        ("request_formal_valuation", {}),  # Not ready at entry: still a usable tool.
        ("propose_facts", {"candidates": candidates[:5]}),
        ("propose_facts", {"candidates": candidates[5:]}),
        ("propose_forecast", forecast(candidates[0]["block_id"])),
        ("finish_response", {"answer": "还可以继续整理更多资料。"}),  # Controller enforces completion.
    ])
    service._clients[session.session_id] = model
    service.store.save_research(session)
    state = service.turn(session.session_id, ResearchTurn(content="开始自动化估值"))
    assert len(model.calls) == 5
    assert not state.get("action")
    assert {f["status"] for f in state["session"]["facts"]} == {"proposed"}
    question = state["session"]["question"]
    assert question["valuation_review"]["baseline_period"] == "2025-12-31"
    assert question["valuation_review"]["historical_periods"] == []
    assert len([e for e in state["events"] if e["type"] == "review.required"]) == 1
    assert state["session"]["forecast_proposal"]["status"] == "proposed"

    approval = (ResearchTurn(content="确认估值方案并计算") if ticker == "600123"
                else ResearchTurn(question_id=question["question_id"], option_id="accept"))
    accepted = service.turn(session.session_id, approval)
    assert accepted["action"] == {"type": "submit_valuation"}
    assert len(model.calls) == 5  # No additional LLM round or second start instruction.
    assert accepted["session"]["forecast_proposal"]["status"] == "confirmed"
    service._clients.pop(session.session_id)  # Calculations and replay need no model/network.
    runner = ValuationRunner(service.store, FinanceTeamModel())
    record = service.submit_valuation(session.session_id, runner)
    result = runner.execute(record.run_id)
    assert str(result.status).startswith("completed"), result.model_dump(mode="json")
    assert result.result.dcf and result.result.sensitivity and result.result.sensitivity_studies
    assert len(result.result.forecast) == 10
    assert result.request.historical_financials == []
    assert result.request.financials.revenue == history()[-1].revenue
    assert "模型推断" in result.request.assumption_evidence["wacc"][0].note
    assert replay_bundle(build_valuation_bundle(service.store, result))["passed"]
    report = ValuationReportExporter()
    assert "预测与FCFF" in load_workbook(BytesIO(report.xlsx(result))).sheetnames
    assert report.pdf(result).startswith(b"%PDF")
    assert service.submit_valuation(session.session_id, runner).run_id == record.run_id


def staged_model(tmp_path, omit=()):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    result = service._facts(session, ProposeFacts(candidates=baseline(service, session, omit=omit)))
    assert result["status"] == "valuation_inputs_staged" and not result.get("_terminal")
    service._propose_forecast(session, ProposeForecast.model_validate(forecast(session.documents[0].file_id + ":1")))
    return service, session


def test_forecast_never_fills_missing_historical_debt_or_approves_facts(tmp_path):
    service, session = staged_model(tmp_path, omit=("interest_bearing_debt",))
    progress = valuation_progress(session, service.valuation_assembler)
    assert not progress["ready_for_review"]
    assert "有息" in progress["blocking_reason"]
    assert all(f.status == "proposed" for f in session.facts)
    assert session.forecast_proposal.status == "proposed"
    assert not service._request_valuation(session).get("_action")
    assert session.question is None


def test_single_pe_method_still_requires_peer_multiples(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.draft.methods = ["pe"]
    session.pending_action = "valuation"
    result = service._facts(
        session,
        ProposeFacts(candidates=baseline(service, session, methods=("pe",))),
    )
    assert result["status"] == "valuation_inputs_staged"
    progress = valuation_progress(session, service.valuation_assembler)
    assert not progress["ready_for_review"]
    assert "PE 当前0家" in progress["blocking_reason"]
    assert "目标公司自身收盘价" in progress["blocking_reason"]
    assert session.question is None


def test_missing_peer_data_does_not_block_a_ready_dcf(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.draft.methods = ["dcf", "pe"]
    session.pending_action = "valuation"
    service._facts(
        session,
        ProposeFacts(candidates=baseline(service, session, methods=("dcf", "pe"))),
    )
    service._propose_forecast(
        session,
        ProposeForecast.model_validate(forecast(session.documents[0].file_id + ":1")),
    )

    progress = valuation_progress(session, service.valuation_assembler)
    assert progress["ready_for_review"] and progress["degraded"]
    assert progress["requested_methods"] == ["dcf", "pe"]
    assert progress["methods"] == ["dcf"]
    assert "pe" in progress["excluded_methods"]

    response = service._request_valuation(session)
    assert response["_terminal"] and session.question.valuation_review["degraded"]
    service.store.save_research(session)
    accepted = service.turn(
        session.session_id,
        ResearchTurn(
            question_id=session.question.question_id,
            option_id="accept",
        ),
    )
    assert accepted["action"] == {"type": "submit_valuation"}
    saved = service.store.get_research(session.session_id)
    assert saved.valuation_methods_override == ["dcf"]
    request = service.valuation_assembler.build(saved)
    assert request.methods == ["dcf"]
    assert request.requested_methods == ["dcf", "pe"]
    assert "pe" in request.excluded_methods
    record = ValuationRunner(service.store, FinanceTeamModel()).run(request)
    assert record.result and record.result.dcf
    assert any("PE因可靠数据不足未进入本次计算" in warning for warning in record.result.warnings)
    exporter = ValuationReportExporter()
    assert exporter.xlsx(record).startswith(b"PK")
    assert exporter.pdf(record).startswith(b"%PDF")


def test_partial_older_history_does_not_block_reviewed_explicit_forecast(tmp_path):
    service, session = staged_model(tmp_path)
    old = session.facts[0].model_copy(deep=True, update={"fact_id": "old_partial", "period": "2023"})
    session.facts.append(old)
    progress = valuation_progress(session, service.valuation_assembler)
    assert progress["ready_for_review"] and progress["historical_periods"] == []
    assert old.status == "proposed"


@pytest.mark.parametrize("change", ["short_path", "missing_scenario", "percent_not_ratio", "nan", "bad_order", "bad_terminal"])
def test_invalid_forecasts_are_rejected_before_staging(change):
    values = forecast()["inputs"]
    if change == "short_path": values["revenue_growth_scenarios"]["base"].pop()
    if change == "missing_scenario": del values["revenue_growth_scenarios"]["pessimistic"]
    if change == "percent_not_ratio": values["revenue_growth_scenarios"]["base"][0] = "8"
    if change == "nan": values["revenue_growth_scenarios"]["base"][0] = "NaN"
    if change == "bad_order": values["revenue_growth_scenarios"]["pessimistic"][0] = "0.20"
    if change == "bad_terminal": values["terminal_growth"] = "0.10"
    with pytest.raises(ValidationError): ForecastInputs.model_validate(values)


def test_stale_plan_cannot_confirm_values_or_forecast_after_scope_changes(tmp_path):
    service, session = staged_model(tmp_path)
    service._request_valuation(session)
    question = session.question.model_copy(deep=True)
    session.draft.company = "另一家公司"
    with pytest.raises(ValueError, match="方案或证据已变化"):
        service._answer_question(session, ResearchTurn(question_id=question.question_id, option_id="accept"))
    assert all(f.status == "proposed" for f in session.facts)
    assert session.forecast_proposal.status == "proposed"


def test_revise_choice_does_not_immediately_reopen_same_review_or_calculate(tmp_path):
    service, session = staged_model(tmp_path)
    service._request_valuation(session)
    service.store.save_research(session)
    state = service.turn(session.session_id, ResearchTurn(question_id=session.question.question_id, option_id="defer"))
    assert state["session"]["question"] is None
    assert not state.get("action")
    assert state["session"]["forecast_proposal"]["status"] == "proposed"


def test_forecast_tool_rejects_unknown_and_search_snippet_evidence(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    source = block("只有搜索摘要", 2, source_type="web_search")
    service.store.save_research_blocks(session.session_id, "file_test", [source])
    for key in ("unknown", source["block_id"]):
        args = {**forecast(), "evidence_ids": [key]}
        with pytest.raises(ValueError, match="引用或搜索摘要"):
            service._propose_forecast(session, ProposeForecast.model_validate(args))
    assert session.forecast_proposal is None


def test_forecast_cannot_start_valuation_for_a_research_only_request(tmp_path):
    service, session, _ = configured_service(tmp_path)
    with pytest.raises(ValueError, match="已请求DCF"):
        service._propose_forecast(session, ProposeForecast.model_validate(forecast()))
    assert not session.pending_action and session.forecast_proposal is None


def test_confirmed_method_change_invalidates_old_forecast_without_blocking_new_method(tmp_path):
    service, session = staged_model(tmp_path)
    old_id = session.forecast_proposal.proposal_id
    revised = session.draft.model_copy(update={"methods": ["pe"]})
    service._question(session, "task", "改用PE？", [("accept", "改用PE"), ("defer", "保留")], proposed_draft=revised)
    service._answer_question(session, ResearchTurn(question_id=session.question.question_id, option_id="accept"))
    assert session.draft.methods == ["pe"] and session.forecast_proposal is None
    event = next(e for e in service.store.list_events(session.session_id) if e.type == "valuation.forecast_invalidated")
    assert event.payload["proposal_id"] == old_id


def test_staged_correction_replaces_only_after_combined_approval(tmp_path):
    service, session, _ = configured_service(tmp_path)
    service._facts(session, ProposeFacts(candidates=baseline(service, session)))
    service._answer_question(session, ResearchTurn(question_id=session.question.question_id, option_id="accept"))
    old = next(f for f in session.facts if f.metric == "common_shares")
    fid = session.documents[0].file_id
    sources = service.store.research_blocks(session.session_id, fid)
    replacement = {"block_id": fid + ":2", "text": f"{session.draft.company} {session.draft.ticker} 2025年合并报表\ncommon_shares（股） 421000000", "location": {}}
    service.store.save_research_blocks(session.session_id, fid, [*sources, replacement])
    session.pending_action = "valuation"
    candidate = item("common_shares", "421000000", unit="股", block_id=fid + ":2", quote="common_shares（股） 421000000")
    result = service._facts(session, ProposeFacts(candidates=[candidate], replaces=[old.fact_id]))
    assert result["status"] == "valuation_inputs_staged"
    assert old.status == "confirmed" and session.facts[-1].status == "proposed"
    service._refresh_pending_candidates(session)
    assert not session.facts[-1].warnings
    service._propose_forecast(session, ProposeForecast.model_validate(forecast(fid + ":1")))
    service._request_valuation(session)
    service.store.save_research(session)
    state = service.turn(session.session_id, ResearchTurn(question_id=session.question.question_id, option_id="accept"))
    assert state["action"] == {"type": "submit_valuation"}
    saved = service.store.get_research(session.session_id)
    assert next(f for f in saved.facts if f.fact_id == old.fact_id).status == "rejected"
    assert service.valuation_assembler.build(saved).financials.common_shares == Decimal(421000000)


def test_future_forecast_evidence_is_rejected(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    source = block("尚未披露的预测依据", published_at="2027-01-01")
    service.store.save_research_blocks(session.session_id, "file_test", [source])
    with pytest.raises(ValueError, match="未来信息"):
        service._propose_forecast(session, ProposeForecast.model_validate(forecast()))


@pytest.mark.parametrize("text,expected", [("开始自动化估值", True), ("开始正式自动估值", True),
    ("不要开始自动化估值", False), ("为什么不开始自动化估值", False)])
def test_automatic_valuation_intent_remains_explicit(text, expected):
    assert requests_valuation(text) is expected


def test_typed_plan_acceptance_requires_visible_full_review_and_no_modifications():
    assert _explicit_fact_acceptance("确认估值方案并计算", valuation_review=True)
    assert not _explicit_fact_acceptance("确认估值方案并计算")
    assert not _explicit_fact_acceptance("确认估值方案并计算，但把WACC改成8%", valuation_review=True)


def test_revising_forecast_over_confirmed_baseline_is_a_review_state_not_a_crash(tmp_path):
    service, session = staged_model(tmp_path)
    service._request_valuation(session)
    service._answer_question(session, ResearchTurn(question_id=session.question.question_id, option_id="accept"))
    args = forecast(session.documents[0].file_id + ":1")
    args["inputs"]["wacc"] = "0.11"
    service._propose_forecast(session, ProposeForecast.model_validate(args))
    prepared = service._prepare(session)
    assert "预测方案尚未确认" in prepared["answer"]
    assert session.status == "collecting"
    assert valuation_progress(session, service.valuation_assembler)["ready_for_review"]
    assert session.forecast_proposal.status == "proposed"
