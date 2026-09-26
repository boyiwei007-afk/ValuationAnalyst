"""Formal, traceable XLSX and PDF exports for completed valuation runs."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path

from valuationagent.schemas.models import RunRecord


BASELINE_FIELDS = (
    ("revenue", "营业收入"),
    ("ebit_margin", "EBIT利润率"),
    ("tax_rate", "有效所得税率"),
    ("depreciation_amortization", "折旧与摊销"),
    ("capital_expenditure", "资本开支"),
    ("change_operating_nwc", "经营性营运资本变动"),
    ("cash_and_non_operating_assets", "现金及非经营性资产"),
    ("interest_bearing_debt", "有息负债"),
    ("common_shares", "普通股股数"),
    ("net_income_parent", "归母净利润"),
    ("ebitda", "EBITDA"),
)


def _number(value):
    return float(value) if isinstance(value, Decimal) else value


def _label(value):
    return {
        "high": "高", "medium": "中", "low": "低", "unknown": "未评估",
        "adequate": "充足", "limited": "有限", "insufficient": "不足",
        "success": "成功", "completed": "完成", "not_available": "不可用",
        "not_applicable": "不适用", "not_applicable_relative_only": "仅相对估值，不使用DCF假设",
        "direct_confirmed_fact": "已确认原始字段", "user": "用户确认", "core": "核心同业",
    }.get(str(value), value)


def _completed(record: RunRecord):
    if record.result is None:
        raise ValueError("估值尚未完成，不能导出正式报告。")
    return record.result


def _baseline_method(snapshot, field):
    if getattr(snapshot, field, snapshot.statement_items.get(field)) is None:
        return "未提供；不用于本次所选估值方法"
    method = snapshot.calculation_methods.get(field, "direct_confirmed_fact")
    reconciliation = snapshot.calculation_methods.get(f"reconciliation.{field}")
    return "；".join(str(_label(part)) for part in (method, reconciliation) if part)


def _source_note(ref):
    return "；".join(part for part in (ref.note, ref.source_url,
        f"SHA-256 {ref.source_sha256}" if ref.source_sha256 else "") if part)


class ValuationReportExporter:
    """Generate reviewer-friendly reports without changing model results."""

    def xlsx(self, record: RunRecord) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, Reference
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ValueError("Excel报告组件未安装，请安装 valuationagent[reports]。") from exc

        result = _completed(record)
        wb = Workbook()
        wb.remove(wb.active)
        navy, teal, pale, amber = "14273D", "0F8B8D", "EAF4F4", "FFF3CD"
        thin = Side(style="thin", color="D7DEE7")

        def sheet(name, widths=(24, 24, 48)):
            ws = wb.create_sheet(name)
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = "A2"
            for index, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(index)].width = width
            return ws

        def header(ws, row=1):
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(bottom=thin)
            ws.row_dimensions[row].height = 28

        def rows(ws, values, *, formula_columns=()):
            for value in values:
                ws.append([_number(item) for item in value])
            for row in ws.iter_rows():
                for cell in row:
                    # Names and source text are data, never spreadsheet code.
                    if cell.data_type == "f" and cell.column not in formula_columns:
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = Border(bottom=thin)

        summary = sheet("估值摘要", (24, 28, 55))
        rows(summary, [
            ["项目", "结果", "说明"],
            ["公司", result.company.name or result.company.ticker, result.company.ticker or ""],
            ["估值基准日", str(result.valuation_date), "所有市场与公开信息不得晚于该日"],
            ["模型版本", result.model_version, f"运行ID：{record.run_id} · 修订v{record.revision}"],
            ["数据质量置信度", result.data_quality.confidence,
             f"可比历史 {result.data_quality.comparable_years}/{result.data_quality.historical_years} 年"],
            ["证据覆盖率", result.data_quality.evidence_coverage,
             "核心字段中带逐项证据的比例"],
            ["可比样本质量", result.data_quality.peer_sample_quality,
             f"行业参数等级 {result.data_quality.industry_parameter_quality} · 元数据 {result.data_quality.industry_metadata_completeness}"],
            ["DCF每股价值", result.dcf.per_share_value if result.dcf else None,
             f"区间 {result.dcf.range_low} – {result.dcf.range_high}" if result.dcf else "未采用"],
            ["退出倍数交叉校验", result.dcf.exit_multiple_per_share if result.dcf else None,
             (
                 f"相对Gordon差异 {result.dcf.terminal_method_gap:.2%}；仅作校验"
                 if result.dcf and result.dcf.terminal_method_gap is not None
                 else "无文档支持的行业倍数或未采用DCF"
             )],
            ["相对估值区间", (
                f"{result.reconciliation.relative_range[0]:.2f} – "
                f"{result.reconciliation.relative_range[1]:.2f}"
                if result.reconciliation.relative_range
                else None
             ), "成功方法的区间并列展示，不强行平均" if result.reconciliation.relative_range else "没有成功方法"],
            ["交叉验证", result.reconciliation.conclusion, "不同方法保持独立，不强行平均"],
            ["结论", result.executive_summary, ""],
            ["原请求估值方法", " / ".join(str(method).upper() for method in record.request.requested_methods),
             "用户在最终方案确认前选择的方法"],
            ["本次实际估值方法", " / ".join(str(method).upper() for method in record.request.methods),
             "仅包含具备可靠输入的方法"],
            *[
                ["未采用方法", method.upper(), reason]
                for method, reason in record.request.excluded_methods.items()
            ],
        ])
        header(summary)
        summary["B6"].number_format = "0.00%"
        summary["B8"].number_format = '¥#,##0.00'

        assumptions = sheet("关键假设", (26, 22, 70))
        assumption_rows = [
            ["参数", "数值", "依据/口径"],
            ["WACC", result.assumptions.wacc if result.dcf else "不适用", result.assumptions.rationale.get("wacc", "")],
            ["永续增长率", result.assumptions.terminal_growth if result.dcf else "不适用", result.assumptions.rationale.get("terminal_growth", "")],
            ["假设来源", result.assumptions.source, ""],
        ]
        assumption_rows += [[f"WACC组成 · {key}", value, ""] for key, value in result.assumptions.wacc_components.items()]
        assumption_rows += [[f"行业参数 · {key}", str(value), ""] for key, value in result.assumptions.industry_parameters.items()]
        assumption_rows += [[f"经营驱动 · {key}", value, ""] for key, value in result.assumptions.operating_drivers.items()]
        assumption_rows += [[f"计算方法 · {key}", value, ""] for key, value in result.assumptions.calculation_methods.items()]
        assumption_rows += [["模型决定", item, ""] for item in result.assumptions.model_decisions]
        rows(assumptions, assumption_rows)
        header(assumptions)
        for cell in (assumptions["B2"], assumptions["B3"]):
            cell.number_format = "0.00%"
            cell.font = Font(color="1F4E78")

        historical = sheet("历史财务", (16, 20, 20, 20, 20, 20, 18, 45, 32))
        rows(historical, [["报告期", "营业收入", "EBIT率", "归母净利润", "EBITDA", "资本开支", "可比状态", "可比说明", "来源"]] + [
            [item.period_end, item.revenue, item.ebit_margin, item.net_income_parent,
             item.ebitda, item.capital_expenditure, item.comparability_status,
             item.comparability_note, item.source_label]
            for item in [*record.request.historical_financials,
                         *([result.effective_financials] if result.effective_financials else [])]
        ])
        header(historical)
        for row in range(2, historical.max_row + 1):
            historical.cell(row, 3).number_format = "0.00%"
            for col in (2, 4, 5, 6):
                historical.cell(row, col).number_format = '#,##0.00'

        derivation = sheet("基期推导", (18, 32, 20, 74, 45, 68))
        derivation_rows = [["性质", "字段", "数值", "计算/取数口径", "证据ID", "原文定位与说明"]]
        base = result.effective_financials
        if base:
            baseline_names = {field for field, _ in BASELINE_FIELDS}
            for field, label in BASELINE_FIELDS:
                refs = base.evidence.get(field, [])
                locations = []
                for ref in refs:
                    location = " · ".join(part for part in [
                        ref.file_id,
                        f"第{ref.page}页" if ref.page else None,
                        f"{ref.sheet}!{ref.cell or ''}" if ref.sheet else ref.cell,
                    ] if part)
                    locations.append("；".join(part for part in [location, ref.note] if part))
                derivation_rows.append([
                    "估值输入",
                    f"{label} ({field})",
                    getattr(base, field),
                    _baseline_method(base, field),
                    "；".join(ref.evidence_id for ref in refs),
                    "\n".join(locations),
                ])
            for field, value in sorted(base.statement_items.items()):
                if field in baseline_names:
                    continue
                derivation_rows.append([
                    "原始/中间科目",
                    field,
                    value,
                    _baseline_method(base, field),
                    "",
                    "用于上述估值输入的确定性复算；逐项证据见对应估值输入。",
                ])
        else:
            derivation_rows.append(["状态", "无有效基期财务", "", "", "", ""])
        rows(derivation, derivation_rows)
        header(derivation)
        for row_index in range(2, derivation.max_row + 1):
            derivation.cell(row_index, 3).number_format = '#,##0.0000'
            derivation.row_dimensions[row_index].height = 66

        forecast = sheet("预测与FCFF", (12, 16, 16, 16, 16, 16, 16, 16, 42, 16, 16, 48))
        rows(forecast, [["年度", "收入增长率", "营业收入", "EBIT率", "EBIT", "NOPAT", "折旧摊销", "FCFF", "审计复算公式", "资本开支", "Δ经营营运资本", "计算口径"]] + [
            [item.year, item.revenue_growth, item.revenue, item.ebit_margin, item.ebit,
             item.nopat, item.depreciation_amortization, item.fcff,
             f"=F{row}+G{row}-J{row}-K{row}", item.capital_expenditure,
             item.change_operating_nwc,
             "；".join(f"{key}={value}" for key, value in item.calculation_methods.items())]
            for row, item in enumerate(result.forecast, 2)
        ], formula_columns=(9,))
        header(forecast)
        for row in range(2, forecast.max_row + 1):
            forecast.cell(row, 2).number_format = "0.00%"
            forecast.cell(row, 4).number_format = "0.00%"
            for col in (3, 5, 6, 7, 8, 10, 11):
                forecast.cell(row, col).number_format = '#,##0.00'
            forecast.cell(row, 9).font = Font(color="008000")
            # Keep the full calculation-method trail visible for audit and replay.
            forecast.row_dimensions[row].height = 76

        dcf_recalc = None
        dcf_data_start = None
        if result.dcf and result.effective_financials and result.forecast:
            dcf_recalc = sheet(
                "DCF复算",
                (12, 16, 20, 16, 20, 14, 20, 20, 20, 20, 20, 16, 16, 20, 16, 22, 22, 28, 28),
            )
            rows(dcf_recalc, [
                ["参数", "数值", "说明"],
                ["WACC", result.assumptions.wacc, "可编辑；修改后公式输出与敏感性表联动"],
                ["永续增长率", result.assumptions.terminal_growth, "可编辑；必须低于WACC"],
                ["现金及非经营性资产", result.effective_financials.cash_and_non_operating_assets, "企业价值到股权价值桥接"],
                ["有息负债", result.effective_financials.interest_bearing_debt, "企业价值到股权价值桥接"],
                ["普通股股数", result.effective_financials.common_shares, "每股价值分母"],
                ["基期营业收入", result.effective_financials.revenue, "最近一期已确认财务事实"],
                ["贴现政策", record.request.discount_policy, "贴现期沿用系统本次运行结果"],
                ["经营必需现金", -result.dcf.bridge.get("operating_cash_requirement", Decimal(0)), "从现金中扣除，仅剩余现金参与股权价值桥接"],
            ])
            header(dcf_recalc)
            for row_index in range(2, 8):
                cell = dcf_recalc.cell(row_index, 2)
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                cell.font = Font(color="0000FF")
            for row_index in (2, 3):
                dcf_recalc.cell(row_index, 2).number_format = "0.00%"
            for row_index in range(4, 8):
                dcf_recalc.cell(row_index, 2).number_format = '#,##0.00'

            dcf_recalc.cell(10, 1, "基准情景DCF复算")
            dcf_recalc.cell(10, 1).font = Font(bold=True, color=navy)
            dcf_headers = [
                "年度", "收入增长率", "营业收入", "EBIT率", "EBIT", "税率", "NOPAT",
                "折旧摊销", "资本开支", "Δ经营营运资本", "FCFF", "贴现期", "贴现因子", "FCFF现值", "现金流比例",
            ]
            dcf_recalc.append(dcf_headers)
            header(dcf_recalc, 11)
            dcf_data_start = 12
            for offset, item in enumerate(result.forecast):
                row_index = dcf_data_start + offset
                previous_revenue = "$B$7" if offset == 0 else f"C{row_index - 1}"
                tax_rate = item.tax_rate
                if tax_rate is None and offset < len(result.assumptions.tax_rate_path):
                    tax_rate = result.assumptions.tax_rate_path[offset]
                if tax_rate is None:
                    tax_rate = result.effective_financials.tax_rate
                dcf_recalc.append([
                    item.year,
                    _number(item.revenue_growth),
                    f"={previous_revenue}*(1+B{row_index})",
                    _number(item.ebit_margin),
                    f"=C{row_index}*D{row_index}",
                    _number(tax_rate),
                    f"=E{row_index}*(1-F{row_index})",
                    _number(item.depreciation_amortization),
                    _number(item.capital_expenditure),
                    _number(item.change_operating_nwc),
                    f"=G{row_index}+H{row_index}-I{row_index}-J{row_index}",
                    _number(item.discount_period),
                    f"=1/(1+$B$2)^L{row_index}",
                    f"=K{row_index}*O{row_index}*M{row_index}",
                    _number(item.cash_flow_fraction),
                ])
                for column in (2, 4, 6, 8, 9, 10, 12):
                    input_cell = dcf_recalc.cell(row_index, column)
                    input_cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    input_cell.font = Font(color="0000FF")
                for column in (2, 4, 6, 13):
                    dcf_recalc.cell(row_index, column).number_format = "0.00%"
                for column in (3, 5, 7, 8, 9, 10, 11, 14):
                    dcf_recalc.cell(row_index, column).number_format = '#,##0.00'
                dcf_recalc.cell(row_index, 12).number_format = "0.00"

            last_row = dcf_data_start + len(result.forecast) - 1
            terminal_period = f"L{last_row}" if record.request.discount_policy == "year_end" else f"(L{last_row}+O{last_row}/2)"
            cash_bridge = "MAX(0,$B$4-$B$9)" if "surplus_cash" in result.dcf.bridge else "$B$4"
            summary_rows = [
                ["公式输出", "数值"],
                ["显性期FCFF现值", f"=SUM(N{dcf_data_start}:N{last_row})"],
                ["终年FCFF", f"=K{last_row}"],
                ["终值", "=Q3*(1+$B$3)/($B$2-$B$3)"],
                ["终值现值", f"=Q4/(1+$B$2)^{terminal_period}"],
                ["企业价值", "=Q2+Q5"],
                ["加：可分配现金及非经营性资产", f"={cash_bridge}"],
                ["减：有息负债", "=$B$5"],
                ["股权价值", "=Q6+Q7-Q8"],
                ["每股价值", "=Q9/$B$6"],
                ["系统本次结果", _number(result.dcf.per_share_value)],
                ["复算差异", "=Q10-Q11"],
                ["复算检查", '=IF(ABS(Q12)<=0.01,"一致","需复核")'],
            ]
            for row_index, values in enumerate(summary_rows, 1):
                dcf_recalc.cell(row_index, 16, values[0])
                dcf_recalc.cell(row_index, 17, values[1])
            header(dcf_recalc, 1)
            for row_index in range(2, 13):
                dcf_recalc.cell(row_index, 17).number_format = '#,##0.00'
            dcf_recalc.freeze_panes = "A12"

            scenario_row = 16
            for column, value in enumerate(["情景", "系统每股价值", "收入增长路径", "EBIT率路径"], 16):
                dcf_recalc.cell(scenario_row, column, value)
            header(dcf_recalc, scenario_row)
            for offset, scenario in enumerate(("pessimistic", "base", "optimistic"), 1):
                output_row = scenario_row + offset
                dcf_recalc.cell(output_row, 16, scenario)
                dcf_recalc.cell(
                    output_row, 17,
                    _number(result.dcf.scenario_values.get(scenario)),
                )
                dcf_recalc.cell(
                    output_row, 18,
                    ", ".join(f"{value:.2%}" for value in result.assumptions.revenue_growth_scenarios.get(scenario, [])),
                )
                dcf_recalc.cell(
                    output_row, 19,
                    ", ".join(f"{value:.2%}" for value in result.assumptions.ebit_margin_scenarios.get(scenario, [])),
                )
                for column in range(16, 20):
                    dcf_recalc.cell(output_row, column).alignment = Alignment(vertical="top", wrap_text=True)
                dcf_recalc.cell(output_row, 17).number_format = '¥#,##0.00'
                dcf_recalc.row_dimensions[output_row].height = 42
            dcf_recalc.column_dimensions["P"].width = 18
            dcf_recalc.column_dimensions["Q"].width = 18
            dcf_recalc.column_dimensions["R"].width = 42
            dcf_recalc.column_dimensions["S"].width = 42

        requested_relative = any(str(method) != "dcf" for method in record.request.methods)
        if requested_relative or result.effective_peers or result.relative:
            relative = sheet("可比公司", (16, 24, 14, 14, 18, 16, 16, 16, 16, 52))
            rows(relative, [["代码", "公司", "P/E", "P/S", "EV/EBITDA", "层级", "筛选得分", "收入增长", "EBIT率", "筛选依据"]] + [
                [peer.ticker, peer.name, peer.pe, peer.ps, peer.ev_ebitda,
                 peer.peer_tier, peer.selection_score, peer.revenue_growth,
                 peer.ebit_margin, peer.rationale]
                for peer in result.effective_peers
            ])
            header(relative)
            for row_index in range(2, relative.max_row + 1):
                for column in (3, 4, 5, 7):
                    relative.cell(row_index, column).number_format = "0.00"
                for column in (8, 9):
                    relative.cell(row_index, column).number_format = "0.00%"
                relative.row_dimensions[row_index].height = 28
            if result.relative:
                relative.append([])
                relative.append(["相对估值结果"])
                title_row = relative.max_row
                relative.cell(title_row, 1).font = Font(bold=True, color=navy)
                relative.append([
                    "方法", "每股价值", "估值区间", "样本数", "样本质量",
                    "异常值数", "统计口径", "纳入公司", "失败原因",
                ])
                result_header = relative.max_row
                header(relative, result_header)
                for item in result.relative:
                    relative.append([
                        item.method.upper(), item.per_share_value,
                        f"{item.range_low:.2f} - {item.range_high:.2f}" if item.range_low is not None and item.range_high is not None else "",
                        item.sample_size, item.sample_quality, item.outlier_count,
                        item.statistic, ", ".join(item.peer_tickers), item.reason or "",
                    ])
                    output_row = relative.max_row
                    relative.cell(output_row, 2).number_format = '¥#,##0.00'
                    relative.row_dimensions[output_row].height = 38

        sensitivity = sheet("敏感性分析", (20, 18, 18, 18, 18, 18, 18, 18, 18, 18, 18, 70))
        growths = sorted({cell.terminal_growth for cell in result.sensitivity})
        waccs = sorted({cell.wacc for cell in result.sensitivity})
        table = [["WACC / 永续增长率", *growths]] if result.sensitivity else [["相对估值敏感性", "见下表：其余条件不变的单因素压力测试"]]
        lookup = {(cell.wacc, cell.terminal_growth): cell for cell in result.sensitivity}
        for wacc in waccs:
            table.append([wacc, *[
                lookup[(wacc, growth)].per_share_value if lookup[(wacc, growth)].valid else "无效"
                for growth in growths
            ]])
        rows(sensitivity, table)
        header(sensitivity)
        if dcf_recalc is not None and dcf_data_start is not None:
            dcf_last_row = dcf_data_start + len(result.forecast) - 1
            for row_index in range(2, 2 + len(waccs)):
                for column_index in range(2, 2 + len(growths)):
                    wacc_ref = f"$A{row_index}"
                    growth_ref = f"{get_column_letter(column_index)}$1"
                    explicit_terms = "+".join(
                        f"'DCF复算'!$K${forecast_row}*'DCF复算'!$O${forecast_row}/(1+{wacc_ref})^'DCF复算'!$L${forecast_row}"
                        for forecast_row in range(dcf_data_start, dcf_last_row + 1)
                    )
                    terminal_period_ref = f"'DCF复算'!$L${dcf_last_row}" if record.request.discount_policy == "year_end" else f"('DCF复算'!$L${dcf_last_row}+'DCF复算'!$O${dcf_last_row}/2)"
                    cash_ref = "MAX(0,'DCF复算'!$B$4-'DCF复算'!$B$9)" if "surplus_cash" in result.dcf.bridge else "'DCF复算'!$B$4"
                    terminal_term = (
                        f"'DCF复算'!$K${dcf_last_row}*(1+{growth_ref})/"
                        f"({wacc_ref}-{growth_ref})/(1+{wacc_ref})^{terminal_period_ref}"
                    )
                    sensitivity.cell(row_index, column_index).value = (
                        f'=IF({wacc_ref}<={growth_ref},"无效",'
                        f"({explicit_terms}+{terminal_term}+{cash_ref}-"
                        f"'DCF复算'!$B$5)/'DCF复算'!$B$6)"
                    )
        for row in range(2, sensitivity.max_row + 1):
            sensitivity.cell(row, 1).number_format = "0.00%"
        for col in range(2, sensitivity.max_column + 1):
            sensitivity.cell(1, col).number_format = "0.00%"
            for row in range(2, sensitivity.max_row + 1):
                sensitivity.cell(row, col).number_format = '¥#,##0.00'

        if result.sensitivity_studies:
            sensitivity.append([])
            sensitivity.append([
                "编号", "参数", "基准输入", "低值输入", "高值输入", "基准每股",
                "低值每股", "高值每股", "最大相对变动", "影响等级", "状态", "说明",
            ])
            detail_header = sensitivity.max_row
            for study in result.sensitivity_studies:
                sensitivity.append([
                    study.study_id, study.parameter, study.baseline_input,
                    study.low_input, study.high_input, study.baseline_per_share,
                    study.low_per_share, study.high_per_share,
                    study.max_relative_change, study.classification,
                    study.status, study.rationale,
                ])
            header(sensitivity, detail_header)
            for row in range(detail_header + 1, sensitivity.max_row + 1):
                for col in (6, 7, 8):
                    sensitivity.cell(row, col).number_format = '¥#,##0.00'
                sensitivity.cell(row, 9).number_format = "0.00%"
            sensitivity.column_dimensions["L"].width = 70

            completed_rows = [
                (study.study_id + " " + study.parameter, study.max_relative_change)
                for study in result.sensitivity_studies
                if study.status == "completed" and study.max_relative_change is not None
            ]
            completed_rows.sort(key=lambda item: item[1], reverse=True)
            if completed_rows:
                chart_start = sensitivity.max_row + 3
                sensitivity.cell(chart_start, 1, "敏感性因素")
                sensitivity.cell(chart_start, 2, "最大相对变动")
                for offset, (label, impact) in enumerate(completed_rows, 1):
                    sensitivity.cell(chart_start + offset, 1, label)
                    sensitivity.cell(chart_start + offset, 2, _number(impact))
                    sensitivity.cell(chart_start + offset, 2).number_format = "0.00%"
                chart = BarChart()
                chart.type = "bar"
                chart.style = 10
                chart.title = "敏感性因素排序"
                chart.y_axis.title = "参数"
                chart.x_axis.title = "相对估值变动"
                chart.x_axis.numFmt = "0%"
                chart.legend = None
                chart.height = 8
                chart.width = 15
                chart.add_data(
                    Reference(sensitivity, min_col=2, min_row=chart_start,
                              max_row=chart_start + len(completed_rows)),
                    titles_from_data=True,
                )
                chart.set_categories(
                    Reference(sensitivity, min_col=1, min_row=chart_start + 1,
                              max_row=chart_start + len(completed_rows))
                )
                sensitivity.add_chart(chart, "N2")

        sources = sheet("来源与风险", (14, 24, 36, 32, 16, 68))
        source_rows = [["性质", "字段/编号", "来源", "定位", "发布日期", "说明"]]
        for field, refs in result.effective_financials.evidence.items() if result.effective_financials else []:
            for ref in refs:
                location = " · ".join(part for part in [
                    ref.file_id,
                    f"第{ref.page}页" if ref.page else None,
                    f"{ref.sheet}!{ref.cell or ''}" if ref.sheet else ref.cell,
                ] if part)
                source_rows.append([
                    "事实", field, ref.source, location,
                    str(ref.published_at) if ref.published_at else "", _source_note(ref),
                ])
        for field, refs in result.assumption_evidence.items():
            for ref in refs:
                location = " · ".join(part for part in [
                    ref.file_id,
                    f"第{ref.page}页" if ref.page else None,
                    f"{ref.sheet}!{ref.cell or ''}" if ref.sheet else ref.cell,
                ] if part)
                source_rows.append([
                    "假设", field, ref.source, location,
                    str(ref.published_at) if ref.published_at else "", _source_note(ref),
                ])
        for peer in result.effective_peers:
            for metric, refs in peer.evidence.items():
                for ref in refs:
                    source_rows.append(["可比事实", f"{peer.ticker} {metric}", ref.source,
                        f"{peer.name} · {peer.as_of_date} · {peer.multiple_basis}", str(ref.published_at or ""), _source_note(ref)])
        source_rows += [
            ["推论", f"Q{index}", "系统质量评估", "", "", note]
            for index, note in enumerate(result.data_quality.notes, 1)
        ]
        source_rows += [["风险", f"W{index}", "系统校验", "", "", warning] for index, warning in enumerate(result.warnings, 1)]
        if not result.warnings:
            source_rows.append(["风险", "", "系统校验", "", "", "本次运行未产生系统警告；仍需人工复核关键假设与同业口径。"])
        rows(sources, source_rows)
        header(sources)
        for row_index in range(2, sources.max_row + 1):
            description = str(sources.cell(row_index, 6).value or "")
            sources.row_dimensions[row_index].height = min(240, max(34, 15 * (2 + len(description) // 40)))

        if not result.forecast:
            wb.remove(forecast)

        for ws in wb.worksheets:
            if ws.title == "DCF复算" and dcf_data_start is not None:
                ws.auto_filter.ref = f"A11:O{dcf_data_start + len(result.forecast) - 1}"
            else:
                ws.auto_filter.ref = ws.dimensions
            for row in range(2, ws.max_row + 1):
                if row % 2 == 0:
                    for cell in ws[row]:
                        if cell.fill.fill_type is None:
                            cell.fill = PatternFill("solid", fgColor=pale)
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        summary["A1"].fill = PatternFill("solid", fgColor=teal)
        sources["A2"].fill = PatternFill("solid", fgColor=amber)
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def pdf(self, record: RunRecord) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFError, TTFont
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import (
                KeepTogether,
                PageBreak,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError as exc:
            raise ValueError("PDF报告组件未安装，请安装 valuationagent[reports]。") from exc

        result = _completed(record)
        font_path = next((path for path in [
            Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")
        ] if path.is_file()), None)
        font = "Helvetica"
        if font_path:
            try:
                pdfmetrics.registerFont(TTFont("ValuationCN", str(font_path), subfontIndex=0))
                font = "ValuationCN"
            except (OSError, TTFError, ValueError):
                font = "Helvetica"
        if font == "Helvetica":
            # Portable CJK fallback; never silently render Chinese as squares.
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font = "STSong-Light"
        styles = getSampleStyleSheet()
        title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName=font,
                               fontSize=22, leading=30, textColor=colors.HexColor("#14273D"),
                               alignment=TA_CENTER, spaceAfter=12)
        h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontName=font,
                            fontSize=14, textColor=colors.HexColor("#0F6F72"), spaceBefore=10, spaceAfter=7)
        body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName=font,
                              fontSize=9, leading=14, textColor=colors.HexColor("#24364B"))
        small = ParagraphStyle("SmallCN", parent=body, fontSize=7.5, leading=11)
        table_header = ParagraphStyle(
            "TableHeaderCN", parent=small, textColor=colors.white
        )

        def p(value, style=body):
            text = str(_label(value) if value is not None else "-")
            text = text.replace("–", "-").replace("—", "-").replace("‑", "-")
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            return Paragraph(text, style)

        def table(data, widths=None, header=True):
            rendered = []
            for row_index, row in enumerate(data):
                style = table_header if header and row_index == 0 else small
                rendered.append([p(cell, style) for cell in row])
            t = Table(rendered, colWidths=widths, repeatRows=1 if header else 0)
            commands = [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D7DEE7")),
                ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [colors.white, colors.HexColor("#F2F7F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
            if header:
                commands += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14273D")),
                             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
            t.setStyle(TableStyle(commands))
            return t

        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm,
                                topMargin=16*mm, bottomMargin=16*mm,
                                title=f"{result.company.name or result.company.ticker}估值报告")

        def page_decor(canvas, document):
            canvas.saveState()
            canvas.setFont(font, 7.5)
            canvas.setFillColor(colors.HexColor("#65758B"))
            canvas.drawString(16*mm, 9*mm, f"ValuationAgent | {result.company.name or result.company.ticker or ''}")
            canvas.drawRightString(A4[0] - 16*mm, 9*mm, f"Page {document.page}")
            canvas.restoreState()

        story = [p("ValuationAgent 正式估值报告", title),
                 p(f"{result.company.name or ''} · {result.company.ticker or ''} · 基准日 {result.valuation_date}"),
                 Spacer(1, 5*mm), p("估值摘要", h2)]
        story.append(table([
            ["方法", "每股价值", "估值区间/状态"],
            ["DCF", result.dcf.per_share_value if result.dcf else "—",
             f"{result.dcf.range_low} – {result.dcf.range_high}" if result.dcf else "未采用"],
            *([["退出倍数校验", result.dcf.exit_multiple_per_share,
                f"相对Gordon差异 {result.dcf.terminal_method_gap:.2%}；不替代主估值"]]
              if result.dcf and result.dcf.exit_multiple_per_share is not None else []),
            *[[item.method.upper(), item.per_share_value or "—",
               f"{item.range_low} – {item.range_high}" if item.status == "success" else item.reason]
              for item in result.relative],
        ], [35*mm, 38*mm, 85*mm]))
        if record.request.excluded_methods:
            story += [p("数据缺失与方法降级披露", h2), table([
                ["原选方法", "处理", "原因"],
                *[
                    [method.upper(), "未进入本次计算", reason]
                    for method, reason in record.request.excluded_methods.items()
                ],
            ], [30*mm, 40*mm, 88*mm])]
        story += [Spacer(1, 3*mm), p(result.executive_summary), p("数据质量", h2), table([
            ["指标", "结论", "说明"],
            ["总体置信度", result.data_quality.confidence,
             f"可比历史 {result.data_quality.comparable_years}/{result.data_quality.historical_years} 年"],
            ["证据覆盖率", f"{result.data_quality.evidence_coverage:.2%}",
             (f"行业参数 {result.data_quality.industry_parameter_quality} · 元数据 {result.data_quality.industry_metadata_completeness}"
              if result.dcf else "仅统计本次相对估值所需的基期财务字段")],
            ["可比样本", result.data_quality.peer_sample_quality,
             "；".join(result.data_quality.notes) or "未产生额外质量提示"],
        ], [38*mm, 34*mm, 86*mm]), p("关键假设", h2), table([
            ["参数", "数值", "依据"],
            ["WACC", f"{result.assumptions.wacc:.2%}" if result.dcf else "不适用", result.assumptions.rationale.get("wacc", "")],
            ["永续增长率", f"{result.assumptions.terminal_growth:.2%}" if result.dcf else "不适用", result.assumptions.rationale.get("terminal_growth", "")],
            ["假设来源", result.assumptions.source, "模型与输入审计轨迹保存在系统中"],
        ], [38*mm, 34*mm, 86*mm])]
        if result.effective_financials:
            base = result.effective_financials
            baseline_rows = [["基期字段", "数值", "计算/取数口径", "证据ID"]]
            for field, label in BASELINE_FIELDS:
                if getattr(base, field) is None:
                    continue
                refs = base.evidence.get(field, [])
                baseline_rows.append([
                    label,
                    getattr(base, field),
                    _baseline_method(base, field),
                    "；".join(ref.evidence_id for ref in refs) or "-",
                ])
            story += [p("基期财务与确定性推导", h2), table(
                baseline_rows,
                [31*mm, 31*mm, 65*mm, 31*mm],
            )]
        if result.forecast:
            story += [PageBreak(), p(f"{len(result.forecast)}年预测与FCFF", h2)]
            story.append(table([["年", "收入增长", "EBIT率", "收入", "FCFF"]] + [
                [item.year, f"{item.revenue_growth:.2%}", f"{item.ebit_margin:.2%}",
                 f"{item.revenue:,.0f}", f"{item.fcff:,.0f}"] for item in result.forecast
            ], [19*mm, 29*mm, 26*mm, 43*mm, 41*mm]))
        if result.dcf:
            story += [p("DCF价值桥与三情景", h2), table([
                ["项目", "数值", "项目", "数值"],
                ["显性期FCFF现值", (
                    f"{result.dcf.present_value_explicit:,.0f}"
                    if result.dcf.present_value_explicit is not None else "-"
                 ), "终值现值", (
                    f"{result.dcf.present_value_terminal:,.0f}"
                    if result.dcf.present_value_terminal is not None else "-"
                 )],
                ["企业价值", f"{result.dcf.enterprise_value:,.0f}",
                 "股权价值", f"{result.dcf.equity_value:,.0f}"],
                ["终值占企业价值", f"{result.dcf.terminal_value_share:.2%}",
                 "基准每股价值", f"{result.dcf.per_share_value:.2f}"],
            ], [42*mm, 37*mm, 42*mm, 37*mm])]
            scenario_labels = {
                "pessimistic": "悲观", "base": "基准", "optimistic": "乐观",
            }
            story.append(table(
                [["情景", "每股价值"]] + [
                    [scenario_labels.get(name, name), f"{value:.2f}"]
                    for name, value in result.dcf.scenario_values.items()
                ],
                [42*mm, 37*mm],
            ))
            story.append(p(
                "公式口径：企业价值 = 显性期FCFF现值 + 终值现值；股权价值按现金、"
                "非经营性资产、经营所需现金和有息负债完成桥接，再除以普通股股数。",
                small,
            ))
        requested_relative = any(str(method) != "dcf" for method in record.request.methods)
        if requested_relative or result.effective_peers or result.relative:
            story += [PageBreak(), p("可比公司与相对估值", h2)]
            if result.effective_peers:
                story.append(table([["代码", "公司", "P/E", "P/S", "EV/EBITDA", "层级/得分"]] + [
                    [peer.ticker, peer.name, peer.pe or "-", peer.ps or "-", peer.ev_ebitda or "-",
                     f"{_label(peer.peer_tier)} / {peer.selection_score if peer.selection_score is not None else '-'}"]
                    for peer in result.effective_peers
                ], [24*mm, 38*mm, 22*mm, 22*mm, 27*mm, 31*mm]))
            story.append(table(
                [["方法", "状态", "每股价值", "区间/原因"]] + [
                    [
                        item.method.upper(), item.status,
                        f"{item.per_share_value:.2f}" if item.per_share_value is not None else "-",
                        (
                            f"{item.range_low:.2f} - {item.range_high:.2f}"
                            if item.status == "success" else item.reason
                        ),
                    ]
                    for item in result.relative
                ],
                [30*mm, 28*mm, 30*mm, 70*mm],
            ))
        else:
            story.append(PageBreak())
        story += [p("敏感性分析", h2)]
        growths = sorted({cell.terminal_growth for cell in result.sensitivity})
        waccs = sorted({cell.wacc for cell in result.sensitivity})
        lookup = {(cell.wacc, cell.terminal_growth): cell for cell in result.sensitivity}
        sensitivity_rows = [["WACC / g", *[f"{g:.2%}" for g in growths]]]
        for wacc in waccs:
            sensitivity_rows.append([f"{wacc:.2%}", *[
                f"{lookup[(wacc, growth)].per_share_value:.2f}" if lookup[(wacc, growth)].valid else "无效"
                for growth in growths]])
        if result.sensitivity:
            story.append(table(sensitivity_rows))
        if result.sensitivity_studies:
            story += [p("敏感性项目与结果", h2)]
            if not result.dcf:
                story.append(p("各项分别测试基期指标或同业倍数上下变动10%，其他条件不变；属于压力测试，不是概率置信区间。", small))
            impact_rows = [["编号", "参数", "低值", "基准", "高值", "影响", "状态"]]
            for study in result.sensitivity_studies:
                impact_rows.append([
                    study.study_id,
                    study.parameter,
                    f"{study.low_per_share:.2f}" if study.low_per_share is not None else "—",
                    f"{study.baseline_per_share:.2f}" if study.baseline_per_share is not None else "—",
                    f"{study.high_per_share:.2f}" if study.high_per_share is not None else "—",
                    (
                        f"{study.max_relative_change:.2%} / {_label(study.classification)}"
                        if study.max_relative_change is not None else "未测试"
                    ),
                    study.status,
                ])
            story.append(table(
                impact_rows,
                [14*mm, 37*mm, 21*mm, 21*mm, 21*mm, 28*mm, 22*mm],
            ))
            unavailable_notes = [
                f"{study.study_id} {study.parameter}：{study.rationale}"
                for study in result.sensitivity_studies
                if study.status == "not_available"
            ]
            if unavailable_notes:
                story.append(KeepTogether([
                    p("未完成敏感性项及原因", h2),
                    p(f"• {unavailable_notes[0]}", small),
                ]))
                story += [p(f"• {note}", small) for note in unavailable_notes[1:]]
        story += [p("风险、限制与追溯", h2)]
        story.append(p(
            "口径说明：已确认财务及其原文定位属于事实；经营路径与资本成本属于假设；"
            "估值区间、敏感性和结论属于模型推论，不构成投资建议。",
            small,
        ))
        notes = result.warnings or ["系统未产生运行警告；关键假设、同业口径与业务判断仍需人工复核。"]
        story += [p(f"• {note}") for note in notes]
        evidence_rows = [["性质", "字段", "来源", "定位"]]
        if result.effective_financials:
            for field, refs in result.effective_financials.evidence.items():
                for ref in refs:
                    location = " · ".join(part for part in [
                        ref.file_id,
                        f"第{ref.page}页" if ref.page else None,
                        f"{ref.sheet}!{ref.cell or ''}" if ref.sheet else ref.cell,
                    ] if part)
                    evidence_rows.append([
                        "事实", field, ref.source,
                        "；".join(part for part in [location, _source_note(ref)] if part),
                    ])
        for field, refs in result.assumption_evidence.items():
            for ref in refs:
                location = " · ".join(part for part in [
                    ref.file_id,
                    f"第{ref.page}页" if ref.page else None,
                    f"{ref.sheet}!{ref.cell or ''}" if ref.sheet else ref.cell,
                ] if part)
                evidence_rows.append([
                    "假设", field, ref.source,
                    "；".join(part for part in [location, _source_note(ref)] if part),
                ])
        for peer in result.effective_peers:
            for metric, refs in peer.evidence.items():
                for ref in refs:
                    evidence_rows.append(["可比事实", f"{peer.ticker} {metric}", ref.source,
                        f"{peer.name} · {peer.as_of_date} · {peer.multiple_basis}；{_source_note(ref)}"])
        if len(evidence_rows) > 1:
            story += [p("证据索引", h2), table(
                evidence_rows,
                [18*mm, 32*mm, 40*mm, 74*mm],
            )]
        story += [Spacer(1, 3*mm), p(
            f"模型版本：{result.model_version}　运行ID：{record.run_id}　输入哈希：{result.effective_input_hash or result.input_hash}", small
        )]
        doc.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
        return output.getvalue()

    def export(self, record: RunRecord, format: str, *, store=None) -> tuple[bytes, str, str]:
        selected = format.lower().lstrip(".")
        if record.result is None and selected in {"pdf", "json"}:
            import json
            from valuationagent.application.result_document import build_run_diagnostic, render_run_diagnostic
            diagnostic = build_run_diagnostic(record, store)
            if selected == "pdf":
                return render_run_diagnostic(diagnostic), "application/pdf", "pdf"
            return json.dumps(diagnostic, ensure_ascii=False, indent=2).encode("utf-8"), "application/json", "json"
        if selected == "xlsx":
            return self.xlsx(record), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
        if selected == "pdf":
            return self.pdf(record), "application/pdf", "pdf"
        if selected == "json":
            import json
            from valuationagent.application.reproducibility import build_valuation_bundle
            return json.dumps(build_valuation_bundle(store, record), ensure_ascii=False, indent=2).encode("utf-8"), "application/json", "json"
        raise ValueError("导出格式仅支持 json、xlsx 或 pdf。")
