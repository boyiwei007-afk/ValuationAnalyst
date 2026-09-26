"""Bounded local parsing. Blocks retain source locations for model citations."""

import csv
import hashlib
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_CHARS = 600_000
MAX_BLOCKS = 1200


def _line_chunks(text, limit=2000):
    """Keep an ordinary statement row intact at the citation boundary.

    A character offset is not a financial row boundary: cutting a label away
    from its amount makes it impossible to verify either part. Exceptionally
    long single lines still obey the hard block limit.
    """
    pending = ""
    for line in text.splitlines(keepends=True):
        if pending and len(pending) + len(line) > limit:
            yield pending
            pending = ""
        while len(line) > limit:
            yield line[:limit]
            line = line[limit:]
        pending += line
    if pending:
        yield pending


class _ReadableHtmlParser(HTMLParser):
    """Extract visible document structure without executing page content."""

    _SKIP = {"script", "style", "noscript", "svg", "canvas", "template"}
    _BREAK = {
        "article", "aside", "blockquote", "div", "footer", "h1", "h2", "h3",
        "h4", "h5", "h6", "header", "li", "main", "p", "section", "table", "tr",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._current: list[str] = []
        self.blocks: list[str] = []

    def _flush(self):
        text = re.sub(r"[ \t\r\f\v]+", " ", "".join(self._current))
        text = re.sub(r"\n\s*\n+", "\n", text).strip(" |\n")
        if text:
            self.blocks.append(text)
        self._current = []

    def handle_starttag(self, tag, attrs):
        tag = tag.casefold()
        if tag in self._SKIP:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag in self._BREAK:
            self._flush()
        elif tag == "br":
            self._current.append("\n")
        elif tag in {"td", "th"} and self._current:
            self._current.append(" | ")

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if not self._skip_depth and tag in self._BREAK:
            self._flush()

    def handle_data(self, data):
        if not self._skip_depth:
            self._current.append(data)

    def close(self):
        super().close()
        self._flush()


def parse_document(meta: dict, *, check_cancel=None, pdf_start_page=1, pdf_page_limit=400, block_offset=0) -> tuple[list[dict], list[str]]:
    path = Path(meta["storage_path"])
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError("文件超过50 MB解析上限，请拆分后上传。")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != meta["sha256"]:
        raise ValueError("文件内容发生变化，请重新上传。")
    blocks, warnings = [], []
    total = 0

    def add(text, location):
        nonlocal total
        if check_cancel:
            check_cancel()
        text = str(text).strip()
        for chunk in _line_chunks(text):
            if total + len(chunk) > MAX_CHARS or len(blocks) >= MAX_BLOCKS:
                if "解析达到大小上限；剩余内容尚未读取。" not in warnings:
                    warnings.append("解析达到大小上限；剩余内容尚未读取。")
                return False
            blocks.append({"block_id": f"{meta['file_id']}:{block_offset + len(blocks) + 1}",
                           "file_id": meta["file_id"], "location": location, "text": chunk})
            total += len(chunk)
        return True

    suffix = path.suffix.lower()
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(i.file_size for i in entries) > 150 * 1024 * 1024:
                raise ValueError("Office文件解压后过大，请拆分相关章节或工作表。")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("PDF 解析依赖未安装，请安装项目 documents 可选依赖。") from None
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("PDF 已加密，请提供可读取的版本。")
        end_page = min(len(reader.pages), pdf_start_page + pdf_page_limit - 1)
        if end_page < len(reader.pages):
            warnings.append(f"本次读取第{pdf_start_page}至{end_page}页（全文{len(reader.pages)}页）；可用read_document的start_page继续读取。")
        for page_no in range(max(1, pdf_start_page), end_page + 1):
            if check_cancel:
                check_cancel()
            page = reader.pages[page_no - 1]
            try:
                # Layout mode retains report columns and table spacing much
                # better than the legacy plain-text order in modern pypdf.
                text = page.extract_text(extraction_mode="layout") or ""
            except (TypeError, ValueError, NotImplementedError):
                text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"第 {page_no} 页未提取到文字；扫描件 OCR 尚未接入。")
            if not add(text, {"page": page_no}):
                break
    elif suffix == ".docx":
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError:
            raise ValueError(
                "DOCX 解析依赖未安装，请安装项目 documents 可选依赖。"
            ) from None
        document = Document(io.BytesIO(raw))
        paragraph_no = table_no = 0
        # python-docx exposes paragraphs and tables as separate lists. Iterating
        # the XML body keeps their original order so a heading remains next to
        # the table it explains, which is essential for reliable extraction.
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph_no += 1
                paragraph = Paragraph(child, document)
                if paragraph.text.strip() and not add(
                    paragraph.text,
                    {
                        "paragraph": paragraph_no,
                        "style": paragraph.style.name,
                    },
                ):
                    break
            elif child.tag.endswith("}tbl"):
                table_no += 1
                table = Table(child, document)
                for row_no, row in enumerate(table.rows, 1):
                    values = [
                        cell.text.replace("\n", " / ").strip() for cell in row.cells
                    ]
                    if any(values) and not add(
                        " | ".join(values),
                        {"table": table_no, "row": row_no},
                    ):
                        break
                if total >= MAX_CHARS or len(blocks) >= MAX_BLOCKS:
                    break
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Excel 解析依赖未安装，请安装项目 documents 可选依赖。") from None
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        cached_book = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
        )
        try:
            for sheet in book.worksheets:
                cached_sheet = cached_book[sheet.title]
                missing_formula_cache = False
                if sheet.max_column and sheet.max_column > 100:
                    warnings.append(f"{sheet.title} 仅读取前 100 列。")
                max_col = min(sheet.max_column or 100, 100)
                # Read-only worksheets stream from their beginning on random
                # access. Zip the two streams once instead of reparsing for
                # every formula cell (quadratic in the number of rows).
                cached_rows = cached_sheet.iter_rows(max_col=max_col)
                for index, row in enumerate(sheet.iter_rows(max_col=max_col), 1):
                    cached_row = next(cached_rows, ())
                    if index > 10000:
                        warnings.append(f"{sheet.title} 仅读取前 10000 行。")
                        break
                    cells = []
                    for column, cell in enumerate(row):
                        if cell.value is not None:
                            if cell.data_type == "f":
                                cached = cached_row[column].value if column < len(cached_row) else None
                                formula = re.sub(
                                    r"(?i)https?://[^\s\"')]+",
                                    "[URL]",
                                    str(cell.value),
                                )
                                if cached is None:
                                    value = f"[公式，未求值: {formula}；缓存值缺失]"
                                    missing_formula_cache = True
                                else:
                                    value = f"[公式: {formula}；缓存值: {cached}]"
                            else:
                                value = str(cell.value)
                            cells.append(f"{cell.coordinate}: {value}")
                    if cells and not add(" | ".join(cells), {"sheet": sheet.title, "row": index}):
                        break
                if total >= MAX_CHARS or len(blocks) >= MAX_BLOCKS:
                    break
                if missing_formula_cache:
                    warnings.append(
                        f"{sheet.title} 含没有缓存结果的公式；系统保留了公式，但未自行执行工作簿计算。"
                    )
        finally:
            book.close()
            cached_book.close()
    elif suffix in {".html", ".htm"}:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
            except UnicodeDecodeError:
                raise ValueError("网页编码无法识别，请另存为 UTF-8。") from None
        parser = _ReadableHtmlParser()
        parser.feed(text)
        parser.close()
        for section, paragraph in enumerate(parser.blocks, 1):
            if not add(paragraph, {"html_section": section}):
                break
        if not blocks:
            warnings.append("网页原文没有可读取的正文，可能依赖脚本渲染或访问权限。")
    elif suffix in {".txt", ".md", ".csv", ".tsv", ".json"}:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
            except UnicodeDecodeError:
                raise ValueError("文本编码无法识别，请另存为 UTF-8。") from None
        if suffix == ".json":
            text = json.dumps(json.loads(text, parse_float=str), ensure_ascii=False, indent=2)
        if suffix in {".csv", ".tsv"}:
            try:
                dialect = csv.Sniffer().sniff(
                    text[:8192], delimiters="\t,;|"
                )
            except csv.Error:
                dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
            for row_no, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                if not add(" | ".join(row), {"row": row_no}):
                    break
        else:
            for line, paragraph in enumerate(re.split(r"\n\s*\n", text), 1):
                if not add(paragraph, {"section": line}):
                    break
    else:
        raise ValueError(
            "此格式尚不能解析；旧版 .xls 请另存为 .xlsx，扫描件请先做 OCR，"
            "或上传 PDF/DOCX/XLSX/CSV/TSV/HTML/JSON/TXT。"
        )
    return blocks, warnings[:30]
